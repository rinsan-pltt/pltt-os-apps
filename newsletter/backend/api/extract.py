"""Extract plain text from uploaded files. Lightweight, dependency-lean
extractors (no docling/model downloads) — PDF, DOCX, PPTX, text/markdown."""

import io


def extract_text(data: bytes, mime: str, filename: str) -> tuple[str, int]:
    """Return (text, page_count). page_count is best-effort (0 if unknown)."""
    name = filename.lower()
    if mime == "application/pdf" or name.endswith(".pdf"):
        return _pdf(data)
    if name.endswith(".docx") or "wordprocessingml" in mime:
        return _docx(data)
    if name.endswith(".pptx") or "presentationml" in mime:
        return _pptx(data)
    # text / markdown / unknown -> decode
    return _text(data)


def _pdf(data: bytes) -> tuple[str, int]:
    import fitz  # pymupdf

    doc = fitz.open(stream=data, filetype="pdf")
    parts = [page.get_text() for page in doc]
    pages = doc.page_count
    doc.close()
    return "\n\n".join(parts).strip(), pages


def _docx(data: bytes) -> tuple[str, int]:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip(), 0


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
    return "\n\n".join(parts).strip(), len(prs.slides)


def _text(data: bytes) -> tuple[str, int]:
    for enc in ("utf-8", "latin-1"):
        try:
            return data.decode(enc).strip(), 1
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore").strip(), 1
