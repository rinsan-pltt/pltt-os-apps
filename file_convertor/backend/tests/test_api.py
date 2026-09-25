"""Smoke tests: health, tool registry, and a real merge/split/compress round-trip."""

from __future__ import annotations

import http.server
import io
import json
import re
import sys
import threading
import zipfile

import fitz
import pytest
from fastapi.testclient import TestClient

from api.core import browser_render, net, webpage_render
from api.main import app

client = TestClient(app)


def make_pdf(pages: int = 3) -> bytes:
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Hello page {i + 1}")
    data = doc.tobytes()
    doc.close()
    return data


def make_text_pdf(text: str) -> bytes:
    """A one-page PDF containing exactly `text`, so a conversion's output can be
    asserted on rather than just its status code."""
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def await_sync(coro):
    """Run one coroutine from a sync test.

    The fake Data Room service is async (it mirrors the platform's surface), but
    the tests around it are sync because TestClient is."""
    import asyncio

    return asyncio.run(coro)


def test_health():
    resp = client.get("/api/health")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}


def test_list_tools():
    resp = client.get("/api/tools")
    assert resp.status_code == 200
    slugs = {t["slug"] for t in resp.json()}
    assert {"merge-pdf", "pdf-to-word", "compress-pdf", "image-converter"} <= slugs


def test_talk_unavailable_standalone():
    # Talk (ctx.talk) only exists inside Palette OS; standalone it 424s cleanly
    # rather than crashing, so the routes are still safe to mount everywhere.
    assert client.get("/api/talk/channels").status_code == 424
    resp = client.post("/api/talk/post", data={"message": "hi", "channel": "general"})
    assert resp.status_code == 424


def test_talk_post_requires_message():
    resp = client.post("/api/talk/post", data={"message": "   ", "channel": "general"})
    assert resp.status_code == 422


def test_merge_pdf():
    a, b = make_pdf(2), make_pdf(3)
    resp = client.post(
        "/api/tools/merge-pdf",
        files=[
            ("files", ("a.pdf", a, "application/pdf")),
            ("files", ("b.pdf", b, "application/pdf")),
        ],
    )
    assert resp.status_code == 200, resp.text
    merged = fitz.open(stream=resp.content, filetype="pdf")
    assert merged.page_count == 5


def test_split_pdf_ranges():
    resp = client.post(
        "/api/tools/split-pdf",
        files=[("files", ("doc.pdf", make_pdf(4), "application/pdf"))],
        data={"ranges": "1-2,4"},
    )
    assert resp.status_code == 200, resp.text
    zf = zipfile.ZipFile(io.BytesIO(resp.content))
    assert len(zf.namelist()) == 2


def test_compress_pdf():
    resp = client.post(
        "/api/tools/compress-pdf",
        files=[("files", ("doc.pdf", make_pdf(3), "application/pdf"))],
        data={"level": "high"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/pdf"


def test_pdf_to_text():
    resp = client.post(
        "/api/tools/pdf-to-text",
        files=[("files", ("doc.pdf", make_pdf(2), "application/pdf"))],
    )
    assert resp.status_code == 200, resp.text
    assert "Hello page 1" in resp.text


def test_protect_then_unlock():
    protected = client.post(
        "/api/tools/protect-pdf",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
        data={"password": "secret123"},
    )
    assert protected.status_code == 200, protected.text

    unlocked = client.post(
        "/api/tools/unlock-pdf",
        files=[("files", ("locked.pdf", protected.content, "application/pdf"))],
        data={"password": "secret123"},
    )
    assert unlocked.status_code == 200, unlocked.text
    doc = fitz.open(stream=unlocked.content, filetype="pdf")
    assert not doc.needs_pass


@pytest.mark.parametrize("password", ["short", "x" * 33])
def test_protect_rejects_password_outside_length_bounds(password):
    resp = client.post(
        "/api/tools/protect-pdf",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
        data={"password": password},
    )
    assert resp.status_code == 422, resp.text


def test_pdf_to_word_basic_fallback(tmp_path):
    """The hosted platform can't install pdf2docx — the fallback must produce
    a readable DOCX from PyMuPDF text extraction alone."""
    from docx import Document

    from api.core.convert_ops import _pdf_to_word_basic

    src = tmp_path / "doc.pdf"
    src.write_bytes(make_pdf(2))
    outputs = _pdf_to_word_basic(src, tmp_path)
    assert len(outputs) == 1 and outputs[0].suffix == ".docx"
    text = "\n".join(p.text for p in Document(str(outputs[0])).paragraphs)
    assert "Hello page 1" in text and "Hello page 2" in text


def test_file_responses_expose_content_disposition_cross_origin():
    """Regression guard: the frontend (:7321) and backend (:8732) are
    different origins under `pltt dev`, so a browser hides
    Content-Disposition from JS on every file download UNLESS the response
    explicitly exposes it via CORS. Without this, every downloaded/chained
    file silently loses its filename+extension in the real app even though
    curl (which isn't CORS-restricted) sees the header fine — exactly what
    broke Workflow chaining (a 'this tool works with: .pdf' 422 on the second
    step, because the first step's output arrived with no extension)."""
    resp = client.post(
        "/api/tools/merge-pdf",
        files=[
            ("files", ("a.pdf", make_pdf(1), "application/pdf")),
            ("files", ("b.pdf", make_pdf(1), "application/pdf")),
        ],
    )
    assert resp.status_code == 200, resp.text
    exposed = resp.headers.get("access-control-expose-headers", "")
    assert "content-disposition" in exposed.lower()


def test_wrong_extension_rejected():
    # Merge takes any document now, so the stranger has to be something that
    # is not a document at all.
    resp = client.post(
        "/api/tools/merge-pdf",
        files=[
            ("files", ("a.zip", b"not a document", "application/zip")),
            ("files", ("b.pdf", make_pdf(1), "application/pdf")),
        ],
    )
    assert resp.status_code == 422


def test_ai_export_txt_and_pdf():
    resp = client.post("/api/ai/export", data={"text": "AI summary text", "format": "txt", "basename": "mydoc"})
    assert resp.status_code == 200, resp.text
    assert resp.text == "AI summary text"
    assert "mydoc.txt" in resp.headers.get("content-disposition", "")

    resp = client.post("/api/ai/export", data={"text": "AI summary text", "format": "pdf", "basename": "mydoc"})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/pdf"


def test_ai_summarize_without_key(monkeypatch):
    monkeypatch.delenv("OPENAI_KEY", raising=False)
    resp = client.post(
        "/api/ai/summarize",
        files={"file": ("notes.txt", b"Some meaningful text to summarize.", "text/plain")},
    )
    assert resp.status_code == 424
    assert "OPENAI_KEY" in resp.json()["detail"]


def test_ai_translate_without_key(monkeypatch):
    monkeypatch.delenv("OPENAI_KEY", raising=False)
    resp = client.post(
        "/api/ai/translate",
        files={"file": ("notes.txt", b"Some meaningful text to translate.", "text/plain")},
        data={"target_language": "French"},
    )
    assert resp.status_code == 424
    assert "OPENAI_KEY" in resp.json()["detail"]


def test_extract_data_images_roundtrip():
    from api.core.html_edit import extract_data_images, restore_data_images

    html = '<p>Hello</p><img src="data:image/png;base64,AAAA=="><p>World</p>'
    stripped, mapping = extract_data_images(html)
    assert "data:image" not in stripped
    assert "__IMG_PLACEHOLDER_0__" in stripped
    assert restore_data_images(stripped, mapping) == html


def test_sign_pdf():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGBA", (300, 100), (0, 0, 0, 0)).save(buf, "PNG")
    resp = client.post(
        "/api/tools/sign-pdf",
        files=[
            ("files", ("doc.pdf", make_pdf(2), "application/pdf")),
            ("files", ("signature.png", buf.getvalue(), "image/png")),
        ],
        data={"page": "first", "position": "bottom-right", "width": "150"},
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    assert doc.page_count == 2
    assert len(doc[0].get_images()) >= 1
    assert len(doc[1].get_images()) == 0  # only the requested page got the signature


def test_sign_pdf_with_xy_position():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGBA", (300, 100), (0, 0, 0, 0)).save(buf, "PNG")
    resp = client.post(
        "/api/tools/sign-pdf",
        files=[
            ("files", ("doc.pdf", make_pdf(2), "application/pdf")),
            ("files", ("signature.png", buf.getvalue(), "image/png")),
        ],
        data={"page": "2", "width": "150", "x": "0.25", "y": "0.6"},
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    assert doc.page_count == 2
    assert len(doc[0].get_images()) == 0
    assert len(doc[1].get_images()) >= 1


def test_sign_pdf_requires_one_pdf_and_one_image():
    resp = client.post(
        "/api/tools/sign-pdf",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
        data={"page": "last"},
    )
    assert resp.status_code == 422


def test_text_extraction_docx(tmp_path):
    from docx import Document

    from api.core.text_extract import extract_text

    src = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Extraction smoke test paragraph.")
    doc.save(str(src))
    assert "Extraction smoke test paragraph." in extract_text(src, tmp_path)


def test_edit_extract_html_from_txt_and_pdf():
    resp = client.post(
        "/api/edit/extract",
        files={"file": ("notes.txt", b"Line one.\nLine two.", "text/plain")},
    )
    assert resp.status_code == 200, resp.text
    pages = resp.json()["pages"]
    assert len(pages) == 1
    assert "<p>Line one.</p>" in pages[0]["html"] and "<p>Line two.</p>" in pages[0]["html"]

    resp = client.post(
        "/api/edit/extract",
        files={"file": ("doc.pdf", make_pdf(2), "application/pdf")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    pages = body["pages"]
    # Each source page comes back as its own chunk, correctly sized.
    assert len(pages) == 2
    # Content must survive whichever PDF pipeline ran (pdf2docx+LibreOffice
    # reconstruction, or the positioned-HTML fallback). LibreOffice may wrap
    # lines, so normalize whitespace before matching.
    assert "Hello page 1" in " ".join(pages[0]["html"].split())
    assert "Hello page 2" in " ".join(pages[1]["html"].split())
    doc = fitz.open(stream=make_pdf(2), filetype="pdf")
    expected_rect = doc[0].rect
    doc.close()
    for page in pages:
        assert page["width_pt"] == pytest.approx(expected_rect.width)
        assert page["height_pt"] == pytest.approx(expected_rect.height)


def test_split_into_pages_positioned_html_fallback():
    # Regression test: PyMuPDF's per-page wrapper closes with `</div>\n</div>`
    # (a newline between the two closing tags), not `</div></div>` -- the
    # splitter must not require them adjacent, or every page collapses into
    # one oversized, wrongly-sized chunk.
    from api.core.html_edit import _pdf_to_positioned_html, split_into_pages

    doc = fitz.open()
    doc.new_page(width=595, height=842).insert_text((72, 72), "Page one")
    doc.new_page(width=595, height=842).insert_text((72, 72), "Page two")
    src_bytes = doc.tobytes()
    doc.close()

    import tempfile
    from pathlib import Path

    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "a4.pdf"
        src.write_bytes(src_bytes)
        html = _pdf_to_positioned_html(src)

    chunks = split_into_pages(html)
    assert len(chunks) == 2
    for chunk in chunks:
        assert chunk.width_pt == pytest.approx(595.0)
        assert chunk.height_pt == pytest.approx(842.0)
        assert chunk.positioned is True
        assert chunk.kind == "page"
    assert "Page one" in chunks[0].html
    assert "Page two" in chunks[1].html
    # `top`/`left` on a paragraph only takes effect once `position: absolute`
    # is set (and its ancestor is `position: relative`) -- without both, the
    # coordinates are silently ignored and everything just flows normally.
    assert "position:absolute" in chunks[0][0]
    assert "position:relative" in html


def _sample_workbook() -> bytes:
    """A small workbook with the things a spreadsheet editor must not lose:
    styled header, a date, a currency amount, a percentage, a formula and a
    merged cell."""
    import datetime as dt

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Q3 Sales"
    for col, header in enumerate(["Region", "Date", "Units", "Price", "Total", "Share"], 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill(fill_type="solid", start_color="FF1F4E79", end_color="FF1F4E79")
        cell.alignment = Alignment(horizontal="center")
    ws.append(["North", dt.date(2026, 7, 3), 120, 19.99, "=C2*D2", 0.284])
    ws["B2"].number_format = "dd mmm yyyy"
    ws["D2"].number_format = '"$"#,##0.00'
    ws["F2"].number_format = "0.0%"
    ws.merge_cells("A4:C4")
    ws["A4"] = "Merged note"
    wb.create_sheet("Notes")["A1"] = "Second sheet"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_edit_extract_spreadsheet_returns_a_cell_grid():
    # A workbook must open as a GRID (column letters, row numbers, cells),
    # not as the flat page render every other format goes through -- that
    # render has no cells to edit at all.
    resp = client.post(
        "/api/edit/extract",
        files={"file": ("sales.xlsx", _sample_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    pages = resp.json()["pages"]
    assert [p["kind"] for p in pages] == ["sheet", "sheet"]
    assert [p["name"] for p in pages] == ["Q3 Sales", "Notes"]
    # A sheet has no paper size: it must not be rendered inside a fixed page.
    assert pages[0]["width_pt"] is None and pages[0]["height_pt"] is None
    assert pages[0]["meta"]["truncated"] is False

    grid = pages[0]["html"]
    assert 'class="xl-colhead"' in grid and 'class="xl-rowhead"' in grid
    assert 'data-cell="A1"' in grid
    # Values are shown the way the workbook formats them, not raw.
    assert ">03 Jul 2026<" in grid
    assert ">$19.99<" in grid
    assert ">28.4%<" in grid
    # ...with the underlying value kept, so saving doesn't stringify them.
    assert 'data-t="dt" data-v="2026-07-03T00:00:00"' in grid
    assert 'data-t="f" data-v="=C2*D2"' in grid
    assert 'colspan="3"' in grid  # the merged A4:C4 note
    assert "font-weight:700" in grid and "background-color:#1F4E79" in grid


def test_edit_export_spreadsheet_round_trip_keeps_cells():
    import html as html_lib

    extracted = client.post(
        "/api/edit/extract",
        files={"file": ("sales.xlsx", _sample_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()["pages"]
    edited = "".join(
        '<div class="fc-page" data-positioned="false" data-kind="sheet" '
        f'data-sheet-name="{html_lib.escape(p["name"], quote=True)}">{p["html"]}</div>'
        for p in extracted
    ).replace(">North<", ">Northwest<")

    resp = client.post(
        "/api/edit/export",
        data={"html": edited, "format": "xlsx", "basename": "sales_edited"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml"
    )

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Q3 Sales", "Notes"]
    ws = wb["Q3 Sales"]
    assert ws["A2"].value == "Northwest"          # the edit came through
    assert ws["B2"].value.date().isoformat() == "2026-07-03"  # still a date
    assert ws["C2"].value == 120                  # still a number
    assert ws["E2"].value == "=C2*D2"             # still a live formula
    assert ws["D2"].number_format == '"$"#,##0.00'
    assert ws["A1"].font.bold is True
    assert [str(r) for r in ws.merged_cells.ranges] == ["A4:C4"]
    # Only real content is written -- not the blank cells the grid pads with.
    assert ws.max_row == 4 and ws.max_column == 6


def test_edit_export_spreadsheet_keeps_browser_normalised_styles():
    # Whatever the grid ships, the moment a cell's style is touched in the
    # editor (any toolbox button) the browser reserialises the declaration:
    # `#1F4E79` comes back as `rgb(31, 78, 121)`. Export has to understand
    # that spelling, or every colour a user sets is silently dropped.
    edited = (
        '<div class="fc-page" data-positioned="false" data-kind="sheet" data-sheet-name="Styled">'
        '<table class="xl-grid"><tbody>'
        '<tr><th class="xl-rowhead">1</th>'
        '<td data-cell="A1" style="font-weight: 700; color: rgb(255, 255, 255); '
        'background-color: rgb(31, 78, 121); text-align: center; '
        'border-bottom: 2px solid rgb(0, 0, 0);">Header</td>'
        '<td data-cell="B1" style="background-color: rgba(0, 0, 0, 0);">Clear</td></tr>'
        "</tbody></table></div>"
    )
    resp = client.post(
        "/api/edit/export",
        data={"html": edited, "format": "xlsx", "basename": "styled"},
    )
    assert resp.status_code == 200, resp.text

    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(resp.content))["Styled"]
    assert ws["A1"].font.bold is True
    assert ws["A1"].font.color.rgb == "FFFFFFFF"
    assert ws["A1"].fill.start_color.rgb == "FF1F4E79"
    assert ws["A1"].alignment.horizontal == "center"
    # Cell rules survive the round trip rather than being flattened away.
    assert ws["A1"].border.bottom.style == "medium"
    assert ws["A1"].border.bottom.color.rgb == "FF000000"
    # A fully transparent background is no fill, not a black one.
    assert ws["B1"].fill.start_color.rgb in ("00000000", None)


def test_edit_export_spreadsheet_as_csv():
    import html as html_lib

    page = client.post(
        "/api/edit/extract",
        files={"file": ("sales.xlsx", _sample_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()["pages"][0]
    edited = (
        '<div class="fc-page" data-positioned="false" data-kind="sheet" '
        f'data-sheet-name="{html_lib.escape(page["name"], quote=True)}">{page["html"]}</div>'
    )
    resp = client.post(
        "/api/edit/export",
        data={"html": edited, "format": "csv", "basename": "sales_edited"},
    )
    assert resp.status_code == 200, resp.text
    text = resp.content.decode("utf-8-sig")
    assert text.splitlines()[0] == "Region,Date,Units,Price,Total,Share"
    assert "North,03 Jul 2026,120,$19.99" in text
    # No trailing sea of commas from the grid's blank padding columns.
    assert ",,,,,,," not in text


def test_edit_extract_csv_opens_as_a_grid():
    resp = client.post(
        "/api/edit/extract",
        files={"file": ("rows.csv", b"Name,Qty\nWidget,3\n", "text/csv")},
    )
    assert resp.status_code == 200, resp.text
    pages = resp.json()["pages"]
    assert len(pages) == 1 and pages[0]["kind"] == "sheet"
    assert 'data-cell="A1"' in pages[0]["html"] and ">Widget<" in pages[0]["html"]


def test_edit_export_html_and_txt():
    edited = "<p>Edited <b>line</b> one.</p><table><tr><td>cell</td></tr></table>"
    resp = client.post(
        "/api/edit/export",
        data={"html": edited, "format": "txt", "basename": "notes_edited"},
    )
    assert resp.status_code == 200, resp.text
    assert "Edited line one." in resp.text and "cell" in resp.text

    resp = client.post(
        "/api/edit/export",
        data={"html": edited, "format": "html", "basename": "notes_edited"},
    )
    assert resp.status_code == 200, resp.text
    assert "notes_edited.html" in resp.headers.get("content-disposition", "")
    assert "Edited <b>line</b> one." in resp.text


def test_strip_positioned_markup_for_pdf_docx_export():
    # Regression test: LibreOffice's HTML importer (and the fitz.Story
    # fallback) don't support `position: absolute` at all -- an empty div
    # sized to a hairline (a table border/fill line, see
    # `_page_drawings_html`) gets Writer's default paragraph height and
    # renders as a giant bar; positioned <p> text just flows in whatever
    # order Writer's importer happens to pick. `_strip_positioned_markup`
    # must remove the border/fill divs entirely and drop the position so the
    # remaining text still exports as plain, readable flowing content.
    from api.core.html_edit import _strip_positioned_markup

    html = (
        '<div id="page0" style="width:595.0pt;height:842.0pt;position:relative;overflow:hidden">'
        '<div style="position:absolute;left:72.00pt;top:120.00pt;width:100.00pt;height:1.00pt;background-color:rgb(0,0,0);"></div>'
        '<p style="top:60.8pt;left:72.0pt;line-height:14.0pt;position:absolute;margin:0">'
        '<span style="font-family:Arial,sans-serif;font-size:14.0pt;color:#000000">Hello</span></p>'
        "</div>"
    )
    stripped = _strip_positioned_markup(html)
    assert "position:absolute" not in stripped
    assert "position:relative" not in stripped
    assert '<div style="position:absolute;left:72.00pt' not in stripped
    assert "Hello" in stripped


def test_ai_export_plain_docx():
    resp = client.post(
        "/api/ai/export",
        data={"text": "Edited line one.\nEdited line two.", "format": "docx", "basename": "notes_edited"},
    )
    assert resp.status_code == 200, resp.text
    assert "notes_edited.docx" in resp.headers.get("content-disposition", "")
    import io

    from docx import Document

    document = Document(io.BytesIO(resp.content))
    assert [p.text for p in document.paragraphs] == ["Edited line one.", "Edited line two."]


def test_watermark_pdf():
    resp = client.post(
        "/api/tools/watermark-pdf",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
        data={"text": "CONFIDENTIAL", "position": "center", "opacity": "40"},
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    assert "CONFIDENTIAL" in doc[0].get_text()


def test_page_numbers_pdf():
    resp = client.post(
        "/api/tools/page-numbers-pdf",
        files=[("files", ("doc.pdf", make_pdf(3), "application/pdf"))],
        data={"position": "bottom-center", "start": "5"},
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    assert "5" in doc[0].get_text() and "7" in doc[2].get_text()


def test_crop_pdf():
    resp = client.post(
        "/api/tools/crop-pdf",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
        data={"margins": "50,50,50,50"},
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    assert doc[0].rect.width < fitz.paper_rect("a4").width


def test_repair_pdf():
    resp = client.post(
        "/api/tools/repair-pdf",
        files=[("files", ("doc.pdf", make_pdf(2), "application/pdf"))],
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    assert doc.page_count == 2


def test_redact_pdf():
    resp = client.post(
        "/api/tools/redact-pdf",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
        data={"term": "Hello page 1"},
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    assert "Hello page 1" not in doc[0].get_text()

    missing = client.post(
        "/api/tools/redact-pdf",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
        data={"term": "not in the document"},
    )
    assert missing.status_code == 422


def test_pdf_to_pdfa():
    resp = client.post(
        "/api/tools/pdf-to-pdfa",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/pdf"


def test_pdf_to_markdown():
    resp = client.post(
        "/api/tools/pdf-to-markdown",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
    )
    assert resp.status_code == 200, resp.text
    # The reconstruction pipeline (pdf2docx + LibreOffice) may wrap short
    # lines with real newlines — normalize whitespace before matching, as in
    # the analogous edit-extract test.
    assert "Hello page 1" in " ".join(resp.text.split())


def make_hwpx(text: str = "Hello from HWPX.") -> bytes:
    import hwpx

    doc = hwpx.HwpxDocument.new()
    doc.add_paragraph(text)
    try:
        return doc.to_bytes()
    finally:
        doc.close()


def test_pdf_to_hwp():
    resp = client.post(
        "/api/tools/pdf-to-hwp",
        files=[("files", ("doc.pdf", make_pdf(1), "application/pdf"))],
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-disposition"].endswith('.hwpx"')

    import hwpx

    doc = hwpx.HwpxDocument.open(resp.content)
    try:
        normalized = " ".join(doc.text.plain().split())
    finally:
        doc.close()
    assert "Hello page 1" in normalized


def test_edit_extract_and_export_hwpx():
    """The document editor opens .hwpx and downloads it back as .hwpx —
    LibreOffice can't import the container, so this rides python-hwpx."""
    resp = client.post(
        "/api/edit/extract",
        files={"file": ("doc.hwpx", make_hwpx("Editable HWPX paragraph."), "application/octet-stream")},
    )
    assert resp.status_code == 200, resp.text
    pages = resp.json()["pages"]
    assert len(pages) == 1
    html = pages[0]["html"]
    # Body content only — a whole <html> document would land inside the
    # editor's contenteditable page.
    assert "<html" not in html.lower() and "<body" not in html.lower()
    assert "Editable HWPX paragraph." in html

    edited = html.replace("Editable HWPX paragraph.", "Edited HWPX paragraph.")
    resp = client.post(
        "/api/edit/export",
        data={"html": f'<div class="fc-page" data-positioned="false">{edited}</div>',
              "format": "hwpx", "basename": "doc_edited"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-disposition"].endswith('.hwpx"')

    import hwpx

    doc = hwpx.HwpxDocument.open(resp.content)
    try:
        assert "Edited HWPX paragraph." in " ".join(doc.text.plain().split())
    finally:
        doc.close()


def test_edit_export_hwp_downloads_as_hwpx():
    """Nothing can write the legacy binary .hwp, so a .hwp original comes back
    as .hwpx rather than silently falling through to a PDF."""
    resp = client.post(
        "/api/edit/export",
        data={"html": "<p>Legacy original.</p>", "format": "hwp", "basename": "legacy_edited"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-disposition"].endswith('.hwpx"')


def test_pdf_to_hwp_keeps_the_source_page_size():
    """A Letter/landscape PDF must not come back re-paginated on the template's
    A4 sheet — the .hwpx section carries the source page size in HWPUNIT."""
    doc = fitz.open()
    page = doc.new_page(width=792, height=612)  # US Letter, landscape
    page.insert_text((90, 90), "Landscape page")
    src = doc.tobytes()
    doc.close()

    resp = client.post(
        "/api/tools/pdf-to-hwp",
        files=[("files", ("wide.pdf", src, "application/pdf"))],
    )
    assert resp.status_code == 200, resp.text

    section = zipfile.ZipFile(io.BytesIO(resp.content)).read("Contents/section0.xml").decode()
    page_pr = re.search(r"<hp:pagePr\b[^>]*>", section)
    assert page_pr is not None, section[:400]
    # HWPUNIT is 1/7200 inch: 792pt x 612pt -> 79200 x 61200.
    assert 'width="79200"' in page_pr.group(0)
    assert 'height="61200"' in page_pr.group(0)
    assert 'landscape="WIDELY"' in page_pr.group(0)


def test_html_to_hwpx_keeps_table_geometry_and_alignment(tmp_path):
    """Column widths, merged cells, row heights and per-cell alignment all
    survive the HTML -> HWPX writer (they used to be flattened to an evenly
    split grid of left-aligned cells)."""
    from api.core.hwpx_ops import html_to_hwpx

    html = (
        '<p align="center">Invoice</p>'
        '<table width="480"><col width="240"/><col width="120"/><col width="120"/>'
        '<tr height="20"><td colspan="3" align="center">Totals</td></tr>'
        '<tr height="18"><td rowspan="2">Widget</td>'
        '<td align="right">2</td><td align="right">1,200</td></tr>'
        '<tr height="18"><td align="right">3</td><td align="right">1,800</td></tr>'
        "</table>"
    )
    out = html_to_hwpx(html, "invoice", tmp_path)
    archive = zipfile.ZipFile(out)
    section = archive.read("Contents/section0.xml").decode()

    assert 'rowCnt="3" colCnt="3"' in section
    cells = re.findall(
        r'<hp:cellAddr colAddr="(\d+)" rowAddr="(\d+)"/>'
        r'<hp:cellSpan colSpan="(\d+)" rowSpan="(\d+)"/>'
        r'<hp:cellSz width="(\d+)" height="(\d+)"',
        section,
    )
    geometry = {(int(col), int(row)): tuple(int(v) for v in rest) for col, row, *rest in cells}
    # 1 px = 75 HWPUNIT: the 240/120/120 px columns keep their 2:1:1 ratio, and
    # each span covers the full width/height of the cells it replaced.
    assert geometry[(0, 0)] == (3, 1, 480 * 75, 20 * 75)
    assert geometry[(0, 1)] == (1, 2, 240 * 75, 2 * 18 * 75)
    assert geometry[(1, 1)][2:] == (120 * 75, 18 * 75)

    # The right-aligned cells reference a paragraph property that is RIGHT.
    header = archive.read("Contents/header.xml").decode()
    right_ids = {
        match.group(1)
        for match in re.finditer(
            r'<hh:paraPr\b[^>]*\bid="(\d+)"[^>]*>\s*<hh:align horizontal="RIGHT"', header
        )
    }
    assert right_ids, "no RIGHT paragraph property was created"
    assert any(f'paraPrIDRef="{para_pr_id}"' in section for para_pr_id in right_ids)


def test_html_to_hwpx_keeps_page_breaks(tmp_path):
    """Each source page boundary (pdf2docx marks them with a hard break) has to
    survive, or a two-page document flows onto a different number of pages."""
    import zipfile

    from api.core.hwpx_ops import html_to_hwpx

    html = (
        "<p>Page one body</p>"
        '<p style="margin-bottom: 0.99cm; page-break-before: always"><br/><br/></p>'
        "<p>Page two body</p>"
    )
    archive = zipfile.ZipFile(html_to_hwpx(html, "paged", tmp_path))
    section = archive.read("Contents/section0.xml").decode()
    header = archive.read("Contents/header.xml").decode()

    break_ids = {
        match.group(1)
        for match in re.finditer(
            r'<hh:paraPr\b[^>]*\bid="(\d+)"[^>]*>(?:(?!</hh:paraPr>).)*?pageBreakBefore="1"',
            header,
            re.S,
        )
    }
    assert break_ids, "no page-break paragraph property was created"
    # The break moves onto the paragraph that actually starts the new page --
    # the blank spacer that carried it is dropped, since the top margin already
    # accounts for that whitespace.
    paragraphs = [
        (match.group(1), " ".join(re.findall(r"<hp:t>([^<]*)</hp:t>", match.group(2))))
        for match in re.finditer(
            r'<hp:p\b[^>]*paraPrIDRef="(\d+)"[^>]*>((?:(?!<hp:p\b).)*?)</hp:p>', section, re.S
        )
    ]
    broken = [text for para_pr_id, text in paragraphs if para_pr_id in break_ids]
    assert broken == ["Page two body"], paragraphs


def test_hwp_to_pdf():
    resp = client.post(
        "/api/tools/hwp-to-pdf",
        files=[("files", ("doc.hwpx", make_hwpx("Hello from HWPX."), "application/octet-stream"))],
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    try:
        text = " ".join(doc[0].get_text().split())
    finally:
        doc.close()
    assert "Hello from HWPX." in text


def test_compare_pdf():
    a = make_pdf(1)
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Hello page 1 changed")
    b = doc.tobytes()
    doc.close()
    resp = client.post(
        "/api/tools/compare-pdf",
        files=[
            ("files", ("a.pdf", a, "application/pdf")),
            ("files", ("b.pdf", b, "application/pdf")),
        ],
    )
    assert resp.status_code == 200, resp.text
    assert "added" in resp.text and "removed" in resp.text


def test_scan_to_pdf():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (200, 100), (50, 100, 200)).save(buf, "PNG")
    resp = client.post(
        "/api/tools/scan-to-pdf",
        files=[("files", ("photo.png", buf.getvalue(), "image/png"))],
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/pdf"


def test_organize_preview_and_apply():
    pdf_bytes = make_pdf(3)
    preview = client.post(
        "/api/organize/preview",
        files={"file": ("doc.pdf", pdf_bytes, "application/pdf")},
    )
    assert preview.status_code == 200, preview.text
    pages = preview.json()["pages"]
    assert len(pages) == 3 and pages[0]["thumbnail"].startswith("data:image/jpeg;base64,")

    # Reverse the page order and rotate the first output page.
    apply_resp = client.post(
        "/api/organize/apply",
        files={"file": ("doc.pdf", pdf_bytes, "application/pdf")},
        data={"order": "3:90,2,1"},
    )
    assert apply_resp.status_code == 200, apply_resp.text
    doc = fitz.open(stream=apply_resp.content, filetype="pdf")
    assert doc.page_count == 3
    assert "Hello page 3" in doc[0].get_text()
    assert doc[0].rotation == 90


def test_organize_apply_bad_index():
    pdf_bytes = make_pdf(2)
    resp = client.post(
        "/api/organize/apply",
        files={"file": ("doc.pdf", pdf_bytes, "application/pdf")},
        data={"order": "1,5"},
    )
    assert resp.status_code == 422


def test_html_to_pdf_blocks_private_urls():
    for bad_url in ("http://127.0.0.1:9999/", "http://169.254.169.254/latest/meta-data/", "ftp://example.com"):
        resp = client.post("/api/tools/html-to-pdf", data={"url": bad_url})
        assert resp.status_code == 422, f"{bad_url} should be rejected"


def test_unknown_tool():
    resp = client.post(
        "/api/tools/does-not-exist",
        files=[("files", ("a.pdf", make_pdf(1), "application/pdf"))],
    )
    assert resp.status_code == 404


# ---------------------------------------------------------------- data room
#
# `ctx.data_rooms` is injected by Palette OS and is absent everywhere else, so
# these drive a fake service with the same async surface — which is exactly what
# palette_sdk's DataRoomsClient docstring says to do. That covers the folder
# structure, the archiving of both directions, and running a tool on a stored
# file; what it cannot cover is a live room, which only exists on the platform.


class FakeDataRoomService:
    """In-memory stand-in for the platform's Data Room client.

    Implements the surface `core/data_room.py` actually calls — `ensure_room`,
    `resolve_folder_path`, `contents`, `upload_file`, `read_file_bytes` — so it
    can be substituted directly. Deliberately does NOT import
    `palette_sdk.DataRoomsClient`: that module only exists when the CLI puts its
    backend-sdk on PYTHONPATH (under `pltt test`), and this suite runs under
    plain pytest, where the whole point is that the platform is absent.
    """

    def __init__(self):
        self.rooms: dict[int, dict] = {}
        self.folders: dict[int, dict] = {}
        self.files: dict[int, dict] = {}
        self.blobs: dict[int, bytes] = {}
        self._next = 1

    def _id(self) -> int:
        self._next += 1
        return self._next

    async def list_rooms(self):
        return list(self.rooms.values())

    async def create_room(self, name, description=None):
        rid = self._id()
        self.rooms[rid] = {"id": rid, "name": name, "description": description}
        return self.rooms[rid]

    async def find_room_by_name(self, name, *, case_sensitive=False):
        for room in self.rooms.values():
            if room["name"] == name:
                return room
        return None

    async def ensure_room(self, name, description=None):
        return await self.find_room_by_name(name) or await self.create_room(name, description)

    async def create_folder(self, room_id, name, parent_folder_id=None):
        fid = self._id()
        self.folders[fid] = {
            "id": fid, "name": name, "data_room_id": room_id, "parent_folder_id": parent_folder_id,
        }
        return self.folders[fid]

    async def find_folder_by_name(self, room_id, name, *, parent_folder_id=None, case_sensitive=False):
        for f in self.folders.values():
            if (
                f["data_room_id"] == room_id
                and f["name"] == name
                and f["parent_folder_id"] == parent_folder_id
            ):
                return f
        return None

    async def ensure_folder(self, room_id, name, parent_folder_id=None):
        found = await self.find_folder_by_name(room_id, name, parent_folder_id=parent_folder_id)
        return found or await self.create_folder(room_id, name, parent_folder_id)

    async def resolve_folder_path(self, room_id, path, *, create=False, case_sensitive=False):
        parts = [path] if isinstance(path, str) else list(path)
        parent = None
        current = None
        for part in parts:
            current = await self.find_folder_by_name(room_id, part, parent_folder_id=parent)
            if current is None:
                if not create:
                    return None
                current = await self.create_folder(room_id, part, parent)
            parent = current["id"]
        return current

    async def contents(self, room_id, folder_id=None):
        return {
            "folders": [
                f for f in self.folders.values()
                if f["data_room_id"] == room_id and f["parent_folder_id"] == folder_id
            ],
            "files": [
                f for f in self.files.values()
                if f["data_room_id"] == room_id and f["folder_id"] == folder_id
            ],
        }

    async def find_file_by_name(self, room_id, name, *, folder_id=None, case_sensitive=False):
        for f in self.files.values():
            if f["data_room_id"] == room_id and f["folder_id"] == folder_id and f["original_filename"] == name:
                return f
        return None

    async def upload_file(self, room_id, filename, content, *, folder_id=None, content_type=None):
        fid = self._id()
        self.files[fid] = {
            "id": fid, "original_filename": filename, "data_room_id": room_id,
            "folder_id": folder_id, "file_url": f"https://example.test/{fid}",
            "file_size": len(content), "mime_type": content_type or "application/octet-stream",
        }
        self.blobs[fid] = content
        return self.files[fid]

    async def read_file_bytes(self, file_id):
        return self.blobs[int(file_id)]

    # --- helpers for the assertions -------------------------------------
    def names_in(self, folder_name: str) -> list[str]:
        folder = next((f for f in self.folders.values() if f["name"] == folder_name), None)
        if folder is None:
            return []
        return sorted(
            f["original_filename"] for f in self.files.values() if f["folder_id"] == folder["id"]
        )


@pytest.fixture
def data_room(monkeypatch):
    """Inject the fake service into every request's context.

    The routes reach the service through `ctx.data_rooms`, and standalone there
    is no ctx at all (`palette.ctx_dependency` is None, so `ctx` arrives as
    None). Overriding the module's `_client`/`available` is what lets the same
    route code be exercised without a platform — everything above those two
    seams is the real code path.
    """
    from file_convertor_backend.api.core import data_room as dr

    service = FakeDataRoomService()
    monkeypatch.setattr(dr, "_client", lambda _ctx: service)
    monkeypatch.setattr(dr, "available", lambda _ctx: True)
    return service


def test_data_room_reports_unavailable_without_the_platform():
    """The plain simulator has no Data Room; that is a state, not an error."""
    resp = client.get("/api/data-room")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["available"] is False
    assert "Palette OS" in body["detail"]
    # The page still gets a shape it can render.
    assert body["uploads"]["files"] == [] and body["results"]["files"] == []


def test_data_room_creates_the_app_folder_structure(data_room):
    resp = client.get("/api/data-room")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["available"] is True
    assert body["app_folder"] == "Document Toolbox"
    assert body["uploads"]["name"] == "Uploads"
    assert body["results"]["name"] == "Results"
    # Uploads/Results are nested under the app folder, not at the room root —
    # a room shared with other apps has to stay legible.
    app = next(f for f in data_room.folders.values() if f["name"] == "Document Toolbox")
    assert app["parent_folder_id"] is None
    for name in ("Uploads", "Results"):
        sub = next(f for f in data_room.folders.values() if f["name"] == name)
        assert sub["parent_folder_id"] == app["id"]


def test_data_room_folder_creation_is_idempotent(data_room):
    for _ in range(3):
        assert client.get("/api/data-room").status_code == 200
    assert len([f for f in data_room.folders.values() if f["name"] == "Uploads"]) == 1
    assert len([f for f in data_room.folders.values() if f["name"] == "Document Toolbox"]) == 1


def test_running_a_tool_archives_both_the_input_and_the_output(data_room):
    resp = client.post(
        "/api/tools/pdf-to-text",
        files=[("files", ("notes.pdf", make_text_pdf("Archived text"), "application/pdf"))],
    )
    assert resp.status_code == 200, resp.text
    assert data_room.names_in("Uploads") == ["notes.pdf"]
    # Results are filed under the tool that produced them, so nothing lands
    # loose in Results itself any more.
    assert data_room.names_in("Results") == []
    # The result is whatever the tool named it; assert on the extension.
    results = data_room.names_in("Document to Text")
    assert len(results) == 1 and results[0].endswith(".txt")

    # …and the listing exposes that folder so the page can offer it.
    listing = client.get("/api/data-room").json()
    tool_folders = {f["name"]: f for f in listing["results"]["folders"]}
    assert "Document to Text" in tool_folders, tool_folders
    assert [f["original_filename"] for f in tool_folders["Document to Text"]["files"]] == results
    assert listing["results"]["files"] == []


def test_a_tool_can_run_on_a_file_already_in_the_data_room(data_room):
    # Put a PDF in Uploads the way an earlier conversion would have.
    listing = client.get("/api/data-room").json()
    assert listing["available"] is True
    uploads_folder = next(f for f in data_room.folders.values() if f["name"] == "Uploads")
    stored = await_sync(
        data_room.upload_file(
            listing["room"]["id"], "stored.pdf", make_text_pdf("From the data room"),
            folder_id=uploads_folder["id"], content_type="application/pdf",
        )
    )

    resp = client.post("/api/tools/pdf-to-text", data={"data_room_file_ids": str(stored["id"])})
    assert resp.status_code == 200, resp.text
    assert "From the data room" in resp.content.decode("utf-8", "replace")
    # No upload was sent, so nothing new should appear in Uploads beyond the
    # file that was already there... but the input IS re-archived by name.
    assert "stored.pdf" in data_room.names_in("Uploads")


def test_a_data_room_file_of_the_wrong_type_is_refused_like_an_upload(data_room):
    listing = client.get("/api/data-room").json()
    uploads_folder = next(f for f in data_room.folders.values() if f["name"] == "Uploads")
    stored = await_sync(
        data_room.upload_file(
            listing["room"]["id"], "notes.txt", b"plain text",
            folder_id=uploads_folder["id"], content_type="text/plain",
        )
    )
    # Compress is PDF-only; the document-wide tools would take a .txt.
    resp = client.post("/api/tools/compress-pdf", data={"data_room_file_ids": str(stored["id"])})
    assert resp.status_code == 422
    assert "can't be used here" in resp.json()["detail"]


def test_each_tool_gets_its_own_results_folder(data_room):
    """Two tools, two folders — the point of grouping."""
    pdf = make_text_pdf("Grouped")
    assert client.post("/api/tools/pdf-to-text",
                       files=[("files", ("a.pdf", pdf, "application/pdf"))]).status_code == 200
    assert client.post("/api/tools/compress-pdf",
                       files=[("files", ("b.pdf", pdf, "application/pdf"))]).status_code == 200

    listing = client.get("/api/data-room").json()
    names = sorted(f["name"] for f in listing["results"]["folders"])
    assert names == ["Compress PDF", "Document to Text"], names
    # every folder actually holds its own output, and none of it leaked loose
    assert all(f["files"] for f in listing["results"]["folders"])
    assert listing["results"]["files"] == []


def test_a_tool_can_run_on_a_result_filed_under_another_tool(data_room):
    """The grouping must not make generated files unreachable as inputs.

    `read_file` used to scan only the two top-level folders, so a result now
    living a level deeper would 404 the moment someone tried to chain it.
    """
    made = client.post(
        "/api/tools/compress-pdf",
        files=[("files", ("chained.pdf", make_text_pdf("Chain me"), "application/pdf"))],
    )
    assert made.status_code == 200, made.text

    listing = client.get("/api/data-room").json()
    folder = next(f for f in listing["results"]["folders"] if f["name"] == "Compress PDF")
    result_id = folder["files"][0]["id"]

    resp = client.post("/api/tools/pdf-to-text", data={"data_room_file_ids": str(result_id)})
    assert resp.status_code == 200, resp.text
    assert "Chain me" in resp.content.decode("utf-8", "replace")


def test_a_tool_title_that_is_not_a_legal_folder_name_is_sanitised(data_room):
    """"PDF to PDF/A" would otherwise put a slash inside a path segment."""
    from file_convertor_backend.api.core import data_room as dr

    assert dr.folder_name_for_tool("PDF to PDF/A") == "PDF to PDF-A"
    assert dr.folder_name_for_tool("") == "Other"


def test_an_unknown_data_room_file_id_is_a_404(data_room):
    resp = client.post("/api/tools/pdf-to-text", data={"data_room_file_ids": "999999"})
    assert resp.status_code == 404
    assert "Document Toolbox" in resp.json()["detail"]


def test_a_malformed_data_room_file_id_is_a_422(data_room):
    resp = client.post("/api/tools/pdf-to-text", data={"data_room_file_ids": "not-a-number"})
    assert resp.status_code == 422
    assert "not a file id" in resp.json()["detail"]


def test_no_input_at_all_is_still_a_422(data_room):
    resp = client.post("/api/tools/pdf-to-text", data={})
    assert resp.status_code == 422
    assert "No files uploaded" in resp.json()["detail"]


def test_archiving_failure_does_not_fail_the_conversion(data_room, monkeypatch):
    """Archiving is a side effect of a job the user asked for.

    If the Data Room rejects the write, the conversion they are waiting on must
    still complete — otherwise a storage hiccup breaks the whole app."""
    from file_convertor_backend.api.core import data_room as dr

    async def boom(*args, **kwargs):
        raise RuntimeError("data room is down")

    monkeypatch.setattr(dr, "save_upload", boom)
    monkeypatch.setattr(dr, "save_result", boom)
    resp = client.post(
        "/api/tools/pdf-to-text",
        files=[("files", ("notes.pdf", make_text_pdf("Still converts"), "application/pdf"))],
    )
    assert resp.status_code == 200, resp.text
    assert "Still converts" in resp.content.decode("utf-8", "replace")


# --------------------------------------------------------------- html-to-pdf
#
# The reported failure: a URL that leads to a "redirecting…" stub produced a
# PDF whose entire content was "Please click here if you are not redirected
# within a few seconds." The fetcher followed HTTP 3xx but not the two
# redirects a browser performs from inside the page.


class _RedirectSite(http.server.BaseHTTPRequestHandler):
    """Serves a meta-refresh stub, a script stub, and the real article."""

    PAGES = {
        "/meta": (
            '<html><head><meta http-equiv="refresh" content="0; url=/article"></head>'
            "<body><p>Please click here if you are not redirected within a few seconds.</p></body></html>"
        ),
        "/js": (
            '<html><body><script>window.location.replace("/article")</script>'
            '<a href="/article">Please click here if you are not redirected.</a></body></html>'
        ),
        "/article": "<html><body><h1>The real article</h1><p>" + ("Body text. " * 60) + "</p></body></html>",
        # A real page that happens to carry a long-delay refresh: must NOT be
        # followed, or a slideshow becomes unconvertible.
        "/slow": (
            '<html><head><meta http-equiv="refresh" content="1800; url=/article"></head>'
            "<body><h1>Slow page</h1><p>" + ("Real content. " * 60) + "</p></body></html>"
        ),
    }

    def do_GET(self):  # noqa: N802 — BaseHTTPRequestHandler's interface
        body = self.PAGES.get(self.path)
        if body is None:
            self.send_error(404)
            return
        payload = body.encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def log_message(self, *args):  # keep the test output clean
        pass


@pytest.fixture
def redirect_site(monkeypatch):
    """A local site, with the SSRF guard relaxed for loopback just for it.

    The guard is the reason this cannot point at 127.0.0.1 unaided, and
    `test_html_to_pdf_blocks_private_urls` still proves the guard is on by
    default.
    """
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _RedirectSite)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    # Every loaded copy of the module, not just the one imported here:
    # api/main.py deliberately re-imports this tree under the unique package
    # name `file_convertor_backend` (so the hosted runtime cannot cross-wire
    # two plugins), which means the routes run a different module object than
    # `from api.core import net` returns. Patching only one leaves the route
    # rejecting 127.0.0.1 while the direct calls sail through.
    for name, module in list(sys.modules.items()):
        if name.endswith("api.core.net") and hasattr(module, "_is_public"):
            monkeypatch.setattr(module, "_is_public", lambda host: True)
    try:
        yield f"http://127.0.0.1:{server.server_port}"
    finally:
        server.shutdown()
        server.server_close()


def test_fetch_follows_meta_refresh_stub(redirect_site):
    page = net.fetch_page(f"{redirect_site}/meta")
    assert page.url.endswith("/article")
    assert "The real article" in page.html
    assert "not redirected" not in page.html


def test_fetch_follows_script_redirect_stub(redirect_site):
    page = net.fetch_page(f"{redirect_site}/js")
    assert page.url.endswith("/article")
    assert "The real article" in page.html


def test_fetch_keeps_a_real_page_with_a_long_refresh(redirect_site):
    page = net.fetch_page(f"{redirect_site}/slow")
    assert page.url.endswith("/slow")
    assert "Slow page" in page.html


def test_html_to_pdf_renders_the_page_behind_a_stub(monkeypatch, redirect_site):
    # Pinned to the no-browser path: this is the meta-refresh following in
    # `core.net`, and a host with Chromium would otherwise exercise the
    # browser's own redirect handling and prove nothing about it.
    for name, module in list(sys.modules.items()):
        if name.endswith("api.routes.webpage"):
            monkeypatch.setattr(module, "browser_available", lambda: False)
    resp = client.post("/api/tools/html-to-pdf", data={"url": f"{redirect_site}/meta"})
    assert resp.status_code == 200, resp.text
    with fitz.open(stream=resp.content, filetype="pdf") as doc:
        text = " ".join(page.get_text() for page in doc)
    assert "The real article" in text
    assert "not redirected" not in text


# The other half of the html-to-pdf report: pages that cannot be scraped at
# all must say so, rather than converting the blocker into a PDF of itself.
# The reported cases were a Google search URL (which answers automated clients
# with its "Enable JavaScript to use search" help page) and a Dribbble shot
# (an empty shell filled in by React).


@pytest.mark.parametrize(
    "text, had_browser, expected_fragment",
    [
        ("Enable JavaScript to use search. Turn on JavaScript to keep searching. " * 4, False, "no browser engine"),
        ("Enable JavaScript to use search. Turn on JavaScript to keep searching. " * 4, True, "refusing automated"),
        ("Our systems have detected unusual traffic from your computer network. " * 3, True, "bot check"),
        ("You must be logged in to see this. " * 8, True, "behind a sign-in"),
        ("", True, "came back empty"),
        ("Loading…", True, "came back empty"),
    ],
)
def test_unconvertible_pages_are_reported(text, had_browser, expected_fragment):
    problem = webpage_render.content_problem(text, had_browser=had_browser)
    assert problem and expected_fragment in problem


@pytest.mark.parametrize(
    "text",
    [
        # example.com in full: a real page that is only 127 characters long.
        "Example Domain This domain is for use in documentation examples without "
        "needing permission. Avoid use in operations. Learn more",
        "An ordinary article with plenty to say. " * 30,
    ],
)
def test_real_pages_are_not_reported_as_problems(text):
    assert webpage_render.content_problem(text, had_browser=True) is None


def test_html_to_pdf_reports_a_blocked_page(monkeypatch, redirect_site):
    """A page that renders but holds only a bot check comes back as a 422 with
    an explanation, not as a 200 with a PDF of the bot check."""
    for name, module in list(sys.modules.items()):
        if name.endswith("api.routes.webpage"):
            monkeypatch.setattr(module, "browser_available", lambda: False)
            blocked = (
                "<html><body><p>Our systems have detected unusual traffic from your "
                "computer network. Please try your request again later.</p></body></html>"
            )
            monkeypatch.setattr(
                module,
                "fetch_page",
                lambda url: net.FetchedPage(url=url, html=blocked),
            )
    resp = client.post("/api/tools/html-to-pdf", data={"url": f"{redirect_site}/article"})
    assert resp.status_code == 422
    assert "bot check" in resp.json()["detail"]


@pytest.mark.skipif(not browser_render.browser_available(), reason="no Chromium on this host")
def test_html_to_pdf_uses_the_browser_when_one_is_installed(redirect_site):
    """With a browser, the same stub converts through Chromium instead."""
    for name, module in list(sys.modules.items()):
        if name.endswith("api.core.browser_render"):
            module._available = True
    resp = client.post("/api/tools/html-to-pdf", data={"url": f"{redirect_site}/meta"})
    assert resp.status_code == 200, resp.text
    with fitz.open(stream=resp.content, filetype="pdf") as doc:
        text = " ".join(page.get_text() for page in doc)
    assert "The real article" in text


# ---------------------------------------------------- document-wide tools
# Plain text stands in for "any document" here: it converts to PDF without
# LibreOffice, so these run on any machine. Word, PowerPoint, Excel and Hancom
# go through the same `documents.ensure_pdf` door.

TEXT_DOC = ("notes.txt", b"Quarterly notes\nSecond line", "text/plain")


def test_document_wide_tool_takes_a_non_pdf_and_returns_a_pdf():
    resp = client.post(
        "/api/tools/rotate-pdf",
        files=[("files", TEXT_DOC)],
        data={"angle": "90"},
    )
    assert resp.status_code == 200, resp.text
    doc = fitz.open(stream=resp.content, filetype="pdf")
    assert doc[0].rotation == 90


def test_merge_mixes_documents_and_pdfs():
    resp = client.post(
        "/api/tools/merge-pdf",
        files=[
            ("files", TEXT_DOC),
            ("files", ("doc.pdf", make_pdf(2), "application/pdf")),
        ],
    )
    assert resp.status_code == 200, resp.text
    assert fitz.open(stream=resp.content, filetype="pdf").page_count == 3


def test_page_preview_accepts_documents():
    resp = client.post("/api/organize/preview", files={"file": TEXT_DOC})
    assert resp.status_code == 200, resp.text
    assert len(resp.json()["pages"]) == 1


def test_compare_pages_accepts_documents():
    resp = client.post("/api/compare/pages", files={"file": TEXT_DOC})
    assert resp.status_code == 200, resp.text
    assert len(resp.json()["pages"]) == 1


@pytest.mark.parametrize("slug", ["compress-pdf", "repair-pdf", "unlock-pdf", "pdf-to-word"])
def test_pdf_only_tools_still_refuse_documents(slug):
    resp = client.post(f"/api/tools/{slug}", files=[("files", TEXT_DOC)], data={"password": "whatever1"})
    assert resp.status_code == 422, resp.text


# ------------------------------------------------------------ Unlock Document

def _locked_docx(text: str, password: str) -> bytes:
    import msoffcrypto.format.ooxml
    from docx import Document

    plain = io.BytesIO()
    doc = Document()
    doc.add_paragraph(text)
    doc.save(plain)
    plain.seek(0)
    locked = io.BytesIO()
    msoffcrypto.format.ooxml.OOXMLFile(plain).encrypt(password, locked)
    return locked.getvalue()


DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_unlock_returns_a_word_file_as_word():
    from docx import Document

    resp = client.post(
        "/api/tools/unlock-pdf",
        files=[("files", ("memo.docx", _locked_docx("Board minutes", "secret123"), DOCX_TYPE))],
        data={"password": "secret123"},
    )
    assert resp.status_code == 200, resp.text
    assert 'memo_unlocked.docx' in resp.headers["content-disposition"]
    assert Document(io.BytesIO(resp.content)).paragraphs[0].text == "Board minutes"


def test_unlock_word_with_the_wrong_password_is_refused():
    resp = client.post(
        "/api/tools/unlock-pdf",
        files=[("files", ("memo.docx", _locked_docx("x", "secret123"), DOCX_TYPE))],
        data={"password": "nope"},
    )
    assert resp.status_code == 422
    assert "Wrong password" in resp.json()["detail"]


def test_unlock_says_so_when_a_word_file_has_no_password():
    from docx import Document

    plain = io.BytesIO()
    Document().save(plain)
    resp = client.post(
        "/api/tools/unlock-pdf",
        files=[("files", ("memo.docx", plain.getvalue(), DOCX_TYPE))],
        data={"password": "whatever"},
    )
    assert resp.status_code == 422
    assert "isn't password-protected" in resp.json()["detail"]


def test_unlock_refuses_formats_it_cannot_decrypt():
    resp = client.post(
        "/api/tools/unlock-pdf",
        files=[("files", ("form.hwpx", b"PK\x03\x04", "application/octet-stream"))],
        data={"password": "whatever"},
    )
    assert resp.status_code == 422


# ----------------------------------------------------------- Protect Document

def test_protect_keeps_a_word_file_as_word_and_unlock_reverses_it():
    import msoffcrypto
    from docx import Document

    plain = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Salary review")
    doc.save(plain)

    locked = client.post(
        "/api/tools/protect-pdf",
        files=[("files", ("memo.docx", plain.getvalue(), DOCX_TYPE))],
        data={"password": "secret123"},
    )
    assert locked.status_code == 200, locked.text
    assert "memo_protected.docx" in locked.headers["content-disposition"]
    assert msoffcrypto.OfficeFile(io.BytesIO(locked.content)).is_encrypted()

    unlocked = client.post(
        "/api/tools/unlock-pdf",
        files=[("files", ("memo_protected.docx", locked.content, DOCX_TYPE))],
        data={"password": "secret123"},
    )
    assert unlocked.status_code == 200, unlocked.text
    assert Document(io.BytesIO(unlocked.content)).paragraphs[0].text == "Salary review"


def test_protect_refuses_a_word_file_that_already_has_a_password():
    resp = client.post(
        "/api/tools/protect-pdf",
        files=[("files", ("memo.docx", _locked_docx("x", "first-pass1"), DOCX_TYPE))],
        data={"password": "secret123"},
    )
    assert resp.status_code == 422
    assert "already has a password" in resp.json()["detail"]


def test_protect_turns_formats_it_cannot_encrypt_into_a_locked_pdf():
    resp = client.post(
        "/api/tools/protect-pdf",
        files=[("files", TEXT_DOC)],
        data={"password": "secret123"},
    )
    assert resp.status_code == 200, resp.text
    assert "notes_protected.pdf" in resp.headers["content-disposition"]
    assert fitz.open(stream=resp.content, filetype="pdf").needs_pass


# --------------------------------------------------------- Edit AI assistant

DOC_BLOCKS = [
    {"id": "b1", "page": 1, "line": 1, "para": 1, "kind": "heading", "size": 20, "bold": True, "text": "Annual report"},
    {"id": "b2", "page": 1, "line": 2, "para": 2, "kind": "line", "size": 11, "text": "Revenue grew strongly."},
]
SHEET_CELLS = [
    {"sheet": "Sales", "ref": "A1", "value": "Item"},
    {"sheet": "Sales", "ref": "C1", "value": "Amount"},
    {"sheet": "Sales", "ref": "C2", "value": "1,200"},
    {"sheet": "Sales", "ref": "C3", "value": "300.5"},
    {"sheet": "Sales", "ref": "C4", "value": "n/a"},
]


def _fake_openai(monkeypatch, reply):
    from file_convertor_backend.api.routes import ai as ai_routes

    seen = {}

    def fake(api_key, messages, *, json_mode=False, model=None):
        seen["messages"] = messages
        seen["json_mode"] = json_mode
        return (reply if isinstance(reply, str) else json.dumps(reply)), "test-model"

    monkeypatch.setenv("OPENAI_KEY", "sk-test")
    monkeypatch.setattr(ai_routes, "call_openai", fake)
    return seen


def test_assistant_without_key(monkeypatch):
    monkeypatch.delenv("OPENAI_KEY", raising=False)
    resp = client.post("/api/ai/assistant", json={"prompt": "hi", "blocks": DOC_BLOCKS})
    assert resp.status_code == 424


def test_assistant_keeps_only_valid_operations(monkeypatch):
    seen = _fake_openai(monkeypatch, {
        "answer": "Made the heading red.",
        "operations": [
            {"op": "style", "ids": ["b1", "b99"], "color": "#FF0000", "bold": "yes"},
            {"op": "style", "ids": ["b2"], "color": "red"},        # not hex -> dropped
            {"op": "align", "ids": ["b2"], "align": "diagonal"},   # bad align -> dropped
            {"op": "replace_text", "id": "b2", "find": "strongly", "replace": "a little"},
            {"op": "delete_everything"},
        ],
    })
    resp = client.post("/api/ai/assistant", json={"prompt": "make the heading red", "blocks": DOC_BLOCKS})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["answer"] == "Made the heading red."
    assert body["operations"] == [
        {"op": "style", "ids": ["b1"], "color": "#FF0000"},
        {"op": "replace_text", "id": "b2", "replace": "a little", "find": "strongly"},
    ]
    assert seen["json_mode"] is True
    assert "b1 p1 line1 para1 heading 20pt bold | Annual report" in seen["messages"][0]["content"]


def test_assistant_does_the_arithmetic_itself(monkeypatch):
    _fake_openai(monkeypatch, {
        "answer": "Total: {{total}} across {{n}} rows, average {{avg}}.",
        "operations": [],
        "calcs": [
            {"name": "total", "op": "sum", "range": "C2:C10", "sheet": "Sales"},
            {"name": "n", "op": "count", "range": "C2:C10"},
            {"name": "avg", "op": "avg", "range": "C2:C3"},
        ],
    })
    resp = client.post("/api/ai/assistant", json={"kind": "sheet", "prompt": "total of C", "cells": SHEET_CELLS})
    assert resp.status_code == 200, resp.text
    assert resp.json()["answer"] == "Total: 1,500.5 across 2 rows, average 750.25."


def test_assistant_expands_cell_ranges(monkeypatch):
    _fake_openai(monkeypatch, {
        "answer": "Header row is bold.",
        "operations": [
            {"op": "cell_style", "sheet": "Sales", "refs": ["A1:C1"], "bold": True},
            {"op": "cell_value", "sheet": "Sales", "ref": "d2", "value": "=C2*2"},
            {"op": "cell_value", "sheet": "Nope", "ref": "A1", "value": "x"},
        ],
    })
    resp = client.post("/api/ai/assistant", json={"kind": "sheet", "prompt": "bold header", "cells": SHEET_CELLS})
    ops = resp.json()["operations"]
    assert ops[0] == {"op": "cell_style", "sheet": "Sales", "refs": ["A1", "B1", "C1"], "bold": True}
    assert ops[1] == {"op": "cell_value", "sheet": "Sales", "ref": "D2", "value": "=C2*2"}
    assert len(ops) == 2


def test_assistant_answers_a_question_without_edits(monkeypatch):
    _fake_openai(monkeypatch, '```json\n{"answer": "It says revenue grew.", "operations": []}\n```')
    resp = client.post("/api/ai/assistant", json={"prompt": "what is line 2?", "blocks": DOC_BLOCKS})
    assert resp.json()["operations"] == []
    assert resp.json()["answer"] == "It says revenue grew."


def test_assistant_unreadable_reply_is_a_502(monkeypatch):
    _fake_openai(monkeypatch, "Sure! I changed it.")
    resp = client.post("/api/ai/assistant", json={"prompt": "x", "blocks": DOC_BLOCKS})
    assert resp.status_code == 502


def test_assistant_can_delete_the_selected_object(monkeypatch):
    seen = _fake_openai(monkeypatch, {
        "answer": "Deleted the selected signature.",
        "operations": [
            {"op": "delete", "ids": ["o2", "b99"]},
            {"op": "delete_selection"},  # no text selection -> dropped
        ],
    })
    objects = [
        {"id": "o1", "page": 1, "kind": "image", "box": [5, 5, 20, 12], "text": "logo"},
        {"id": "o2", "page": 1, "kind": "shape", "box": [70, 85, 95, 95], "text": "Authorised Signatory", "selected": True},
    ]
    resp = client.post(
        "/api/ai/assistant",
        json={"prompt": "delete the one I selected", "blocks": DOC_BLOCKS, "objects": objects},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["operations"] == [{"op": "delete", "ids": ["o2"]}]
    outline = seen["messages"][0]["content"]
    # The selected object is listed first, flagged, so a cap can never cut it.
    assert outline.index("o2 p1 shape box 70,85,95,95% selected | Authorised Signatory") < outline.index("b1 p1")


def test_assistant_selection_ops_need_a_selection(monkeypatch):
    seen = _fake_openai(monkeypatch, {
        "answer": "Replaced it.",
        "operations": [{"op": "replace_selection", "text": "modestly"}],
    })
    resp = client.post(
        "/api/ai/assistant",
        json={"prompt": "change this to modestly", "blocks": DOC_BLOCKS, "selection": "strongly"},
    )
    assert resp.json()["operations"] == [{"op": "replace_selection", "text": "modestly"}]
    assert "SELECTION: strongly" in seen["messages"][0]["content"]


def test_assistant_styles_only_the_selected_text(monkeypatch):
    seen = _fake_openai(monkeypatch, {
        "answer": "Coloured the selected text red.",
        "operations": [{"op": "style_selection", "color": "#FF0000", "size": "huge"}],
    })
    blocks = [dict(DOC_BLOCKS[1], selected=True), DOC_BLOCKS[0]]
    resp = client.post(
        "/api/ai/assistant",
        json={"prompt": "make the selected text red", "blocks": blocks, "selection": "grew"},
    )
    assert resp.json()["operations"] == [{"op": "style_selection", "color": "#FF0000"}]
    assert "never on the whole lines" in seen["messages"][0]["content"]

    # Without a selection there is nothing to style.
    resp = client.post("/api/ai/assistant", json={"prompt": "make it red", "blocks": DOC_BLOCKS})
    assert resp.json()["operations"] == []


# ------------------------------------------------ Correct Mistakes (layout)

def _fake_proofreader(monkeypatch, pick):
    """Fake model: `pick(prompt_text)` returns the fixes list for one batch."""
    from file_convertor_backend.api.routes import ai as ai_routes

    def fake(api_key, messages, *, json_mode=False, model=None):
        assert json_mode
        return json.dumps({"fixes": pick(messages[-1]["content"])}), "test-model"

    monkeypatch.setenv("OPENAI_KEY", "sk-test")
    monkeypatch.setattr(ai_routes, "call_openai", fake)


def _segment_id(prompt: str, needle: str) -> str:
    line = next(l for l in prompt.splitlines() if needle in l)
    return line.split(" | ", 1)[0]


def test_proofread_fixes_text_inside_the_original_pdf_layout(monkeypatch):
    _fake_proofreader(monkeypatch, lambda prompt: [
        {"id": _segment_id(prompt, "Teh"), "before": "Teh", "after": "The", "type": "spelling", "explanation": "Typo."},
        {"id": _segment_id(prompt, "Teh"), "before": "not in the line", "after": "x", "type": "grammar"},
    ])
    resp = client.post(
        "/api/ai/proofread",
        files={"file": ("memo.pdf", make_text_pdf("Teh quick brown fox"), "application/pdf")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    page = body["pages"][0]
    # Still the editor's positioned page: the line keeps its coordinates.
    assert page["positioned"] is True and page["width_pt"]
    assert "The quick brown fox" in page["html"] and "Teh" not in page["html"]
    assert "position:absolute" in page["html"].replace(" ", "")
    assert body["fix_count"] == 1
    assert [f["applied"] for f in body["fixes"]] == [True, False]
    assert body["fixes"][0]["page"] == 1


def test_proofread_replacement_across_styled_spans_keeps_the_spans():
    from lxml import html as lxml_html

    from file_convertor_backend.api.core.proofread import replace_in_element

    p = lxml_html.fragment_fromstring('<p style="top:10pt"><b>recie</b><i>ved it</i></p>')
    assert replace_in_element(p, "recieved", "received")
    out = lxml_html.tostring(p, encoding="unicode")
    assert out == '<p style="top:10pt"><b>received</b><i> it</i></p>'


def test_proofread_leaves_spreadsheet_values_alone(monkeypatch):
    prompts = []

    def pick(prompt):
        prompts.append(prompt)
        return [{"id": _segment_id(prompt, "Merged note"), "before": "Merged note", "after": "Merged notes"}]

    _fake_proofreader(monkeypatch, pick)
    resp = client.post(
        "/api/ai/proofread",
        files={"file": ("sales.xlsx", _sample_workbook(), "application/octet-stream")},
    )
    assert resp.status_code == 200, resp.text
    sent = "\n".join(prompts)
    # Text cells are proofread; dates, prices and formulas are not even sent.
    assert "Region" in sent and "Merged note" in sent
    assert "=C2*D2" not in sent and "19.99" not in sent and "Jul" not in sent
    sheet = resp.json()["pages"][0]
    assert sheet["kind"] == "sheet" and "Merged notes" in sheet["html"]
    assert 'data-v="=C2*D2"' in sheet["html"]


def test_proofread_bad_model_reply_is_a_502(monkeypatch):
    from file_convertor_backend.api.routes import ai as ai_routes

    monkeypatch.setenv("OPENAI_KEY", "sk-test")
    monkeypatch.setattr(ai_routes, "call_openai", lambda *a, **k: ("not json", "m"))
    resp = client.post(
        "/api/ai/proofread",
        files={"file": ("memo.pdf", make_text_pdf("Teh quick brown fox"), "application/pdf")},
    )
    assert resp.status_code == 502


def test_proofread_overlapping_fixes_both_land():
    from lxml import html as lxml_html

    from file_convertor_backend.api.core import proofread as pr

    root = lxml_html.fragment_fromstring("<p>Our team have acheived results.</p>", create_parent="div")
    page = pr._Page(data={"html": "", "positioned": False}, root=root)
    loaded, segments = [page], []
    seg = pr._Segment(id="s1", element=root[0], page=0, text=pr.segment_text(root[0]))
    segments.append(seg)
    out = pr.apply_fixes(loaded, segments, [
        {"id": "s1", "before": "team have", "after": "team has", "type": "grammar", "explanation": ""},
        {"id": "s1", "before": "have acheived", "after": "have achieved", "type": "spelling", "explanation": ""},
    ])
    assert [f["applied"] for f in out] == [True, True]
    assert pr.segment_text(root[0]) == "Our team has achieved results."


# ----------------------------------------------------------- Document chat

def test_document_chat_open_reads_the_file_once(monkeypatch):
    monkeypatch.setenv("OPENAI_KEY", "sk-test")
    resp = client.post(
        "/api/ai/chat/open",
        files={"file": ("memo.pdf", make_text_pdf("Budget is 42 dollars"), "application/pdf")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert "Budget is 42 dollars" in body["text"] and body["filename"] == "memo.pdf"


def test_document_chat_answers_with_the_document_and_history(monkeypatch):
    seen = _fake_openai(monkeypatch, "It is 42 dollars.")
    resp = client.post("/api/ai/chat", json={
        "filename": "memo.pdf",
        "text": "Budget is 42 dollars",
        "history": [{"role": "user", "content": "hi"}, {"role": "assistant", "content": "hello"}],
        "prompt": "What is the budget?",
    })
    assert resp.status_code == 200, resp.text
    assert resp.json()["answer"] == "It is 42 dollars."
    msgs = seen["messages"]
    assert "Budget is 42 dollars" in msgs[0]["content"]
    assert [m["role"] for m in msgs] == ["system", "user", "assistant", "user"]
    assert seen["json_mode"] is False


def test_document_chat_summary_chip_keeps_its_length(monkeypatch):
    seen = _fake_openai(monkeypatch, "Short.")
    client.post("/api/ai/chat", json={"text": "x", "prompt": "Short summary", "summary": "short"})
    assert "3-5 sentences" in seen["messages"][-1]["content"]


def test_document_chat_without_key(monkeypatch):
    monkeypatch.delenv("OPENAI_KEY", raising=False)
    assert client.post("/api/ai/chat", json={"text": "x", "prompt": "q"}).status_code == 424


# ------------------------------------------------------------- Global chat

AGENT_CATALOG = [
    {"slug": "compress-pdf", "title": "Compress PDF", "description": "Smaller PDF", "accept": ".pdf",
     "options": [{"name": "level", "kind": "select", "choices": ["low", "medium", "high"]}]},
    {"slug": "protect-pdf", "title": "Protect Document", "description": "Add a password", "accept": ".pdf,.docx",
     "options": [{"name": "password", "kind": "password", "required": True}]},
    {"slug": "word-to-pdf", "title": "Word to PDF", "description": "DOCX to PDF", "accept": ".docx", "options": []},
    {"slug": "edit-document", "title": "Edit Document", "description": "Editor", "accept": ".pdf", "runnable": False},
]
AGENT_FILES = [{"id": "f1", "name": "report.docx", "origin": "uploaded"},
               {"id": "f2", "name": "old.pdf", "origin": "uploaded", "available": False}]


def test_agent_plan_is_validated(monkeypatch):
    seen = _fake_openai(monkeypatch, {
        "reply": "Converting, then compressing.",
        "steps": [
            {"tool": "word-to-pdf", "files": ["f1"]},
            {"tool": "compress-pdf", "files": ["$prev"], "options": {"level": "high", "junk": "x"}},
            {"tool": "compress-pdf", "files": ["f2"]},                    # file not available
            {"tool": "protect-pdf", "files": ["f1"], "options": {}},       # missing password
            {"tool": "edit-document", "files": ["f1"]},                    # interactive
            {"tool": "nonexistent", "files": ["f1"]},
        ],
        "open_tool": "edit-document",
    })
    resp = client.post("/api/ai/agent", json={
        "catalog": AGENT_CATALOG, "files": AGENT_FILES, "prompt": "make it a small pdf",
    })
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["steps"] == [
        {"tool": "word-to-pdf", "files": ["f1"], "options": {}},
        {"tool": "compress-pdf", "files": ["$prev"], "options": {"level": "high"}},
    ]
    assert body["open_tool"] == "edit-document" and body["dropped"] is True
    system = seen["messages"][0]["content"]
    assert "compress-pdf — Compress PDF" in system and "level=low|medium|high" in system
    assert "f2 — old.pdf — uploaded (no longer available" in system
    assert "edit-document" in system and "[interactive]" in system


def test_agent_first_step_cannot_start_from_prev(monkeypatch):
    _fake_openai(monkeypatch, {"reply": "ok", "steps": [{"tool": "compress-pdf", "files": ["$prev"]}]})
    resp = client.post("/api/ai/agent", json={"catalog": AGENT_CATALOG, "files": AGENT_FILES, "prompt": "x"})
    assert resp.json()["steps"] == []


def test_agent_plain_chat(monkeypatch):
    _fake_openai(monkeypatch, {"reply": "Hello! Attach a file to get started.", "steps": []})
    resp = client.post("/api/ai/agent", json={"catalog": AGENT_CATALOG, "prompt": "hi"})
    assert resp.json()["reply"].startswith("Hello") and resp.json()["steps"] == []
