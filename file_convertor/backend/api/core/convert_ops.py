"""Cross-format conversions: PDF <-> Office, PDF <-> images, image <-> image, text."""

from __future__ import annotations

import shutil
from pathlib import Path

from fastapi import HTTPException

from .files import output_dir
from .office import convert_with_soffice


def _open_fitz(src: Path):
    import fitz

    try:
        doc = fitz.open(str(src))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"'{src.name}' is not a valid PDF ({exc}).")
    if doc.needs_pass:
        raise HTTPException(status_code=422, detail="PDF is password-protected. Unlock it first.")
    if doc.page_count == 0:
        raise HTTPException(status_code=422, detail="PDF has no pages.")
    return doc


# ---------------------------------------------------------------- PDF -> Word
def pdf_to_word(src: Path, workdir: Path) -> list[Path]:
    # pdf2docx gives the best layout fidelity but depends on opencv-python,
    # a large native wheel the hosted platform cannot install. Use it when
    # present (local dev installs it via backend/requirements.txt) and fall
    # back to a PyMuPDF + python-docx text rebuild everywhere else.
    try:
        from pdf2docx import Converter
    except ImportError:
        return _pdf_to_word_basic(src, workdir)

    out = output_dir(workdir) / f"{src.stem}.docx"
    try:
        converter = Converter(str(src))
        converter.convert(str(out))
        converter.close()
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=500, detail=f"PDF to Word conversion failed: {exc}")
    if not out.exists():
        raise HTTPException(status_code=500, detail="PDF to Word conversion produced no file.")
    return [out]


def _pdf_to_word_basic(src: Path, workdir: Path) -> list[Path]:
    """Dependency-light PDF -> DOCX: rebuild paragraphs from PyMuPDF text
    blocks, keeping font size, bold and italic per span. Layout (columns,
    tables, images) is not preserved — but it works on any deployment."""
    import fitz
    from docx import Document
    from docx.shared import Pt

    pdf = _open_fitz(src)
    document = Document()
    for page_no, page in enumerate(pdf):
        if page_no:
            document.add_page_break()
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:  # 0 = text block (1 = image)
                continue
            paragraph = document.add_paragraph()
            lines = block.get("lines", [])
            for line_no, line in enumerate(lines):
                if line_no:
                    paragraph.add_run(" ")
                for span in line.get("spans", []):
                    text = span.get("text", "")
                    if not text:
                        continue
                    run = paragraph.add_run(text)
                    size = round(span.get("size", 11))
                    run.font.size = Pt(max(6, min(size, 72)))
                    flags = span.get("flags", 0)
                    run.bold = bool(flags & 16)
                    run.italic = bool(flags & 2)
    pdf.close()

    out = output_dir(workdir) / f"{src.stem}.docx"
    document.save(str(out))
    return [out]


# ---------------------------------------------------------- PDF -> PowerPoint
def pdf_to_powerpoint(src: Path, workdir: Path) -> list[Path]:
    """Render each PDF page and place it as a full-bleed slide image."""
    import fitz
    from pptx import Presentation
    from pptx.util import Emu

    doc = _open_fitz(src)
    out = output_dir(workdir) / f"{src.stem}.pptx"
    prs = Presentation()

    first = doc[0].rect
    # EMU: 914400 per inch; PDF points: 72 per inch.
    prs.slide_width = Emu(int(first.width / 72 * 914400))
    prs.slide_height = Emu(int(first.height / 72 * 914400))
    blank = prs.slide_layouts[6]

    img_dir = workdir / "slides"
    img_dir.mkdir(exist_ok=True)
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
        img_path = img_dir / f"slide_{i}.png"
        pix.save(str(img_path))
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img_path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    doc.close()
    prs.save(str(out))
    return [out]


# --------------------------------------------------------------- PDF -> Excel
def pdf_to_excel(src: Path, workdir: Path) -> list[Path]:
    """Extract tables per page with pdfplumber; fall back to text lines."""
    import pdfplumber
    from openpyxl import Workbook

    out = output_dir(workdir) / f"{src.stem}.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    wrote_anything = False

    try:
        with pdfplumber.open(str(src)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables()
                ws = wb.create_sheet(title=f"Page {page_no}"[:31])
                row_cursor = 1
                if tables:
                    for table in tables:
                        for row in table:
                            for col, cell in enumerate(row, start=1):
                                ws.cell(row=row_cursor, column=col, value=cell)
                            row_cursor += 1
                        row_cursor += 1  # blank row between tables
                        wrote_anything = True
                else:
                    text = page.extract_text() or ""
                    for line in text.splitlines():
                        ws.cell(row=row_cursor, column=1, value=line)
                        row_cursor += 1
                        wrote_anything = True
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Could not read '{src.name}' ({exc}).")

    if not wb.sheetnames:
        wb.create_sheet(title="Empty")
    if not wrote_anything:
        raise HTTPException(
            status_code=422,
            detail="No tables or text found in this PDF (it may be a scanned document).",
        )
    wb.save(str(out))
    return [out]


# ----------------------------------------------------------------- PDF -> JPG
def pdf_to_jpg(src: Path, workdir: Path, dpi: int = 150) -> list[Path]:
    import fitz

    doc = _open_fitz(src)
    out = output_dir(workdir)
    outputs: list[Path] = []
    zoom = dpi / 72
    for i, page in enumerate(doc, start=1):
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
        path = out / f"{src.stem}_page_{i}.jpg"
        pix.save(str(path), jpg_quality=90)
        outputs.append(path)
    doc.close()
    return outputs


# ---------------------------------------------------------------- PDF -> Text
def pdf_to_text(src: Path, workdir: Path) -> list[Path]:
    doc = _open_fitz(src)
    text = "\n\n".join(page.get_text() for page in doc)
    doc.close()
    out = output_dir(workdir) / f"{src.stem}.txt"
    out.write_text(text, encoding="utf-8")
    return [out]


# -------------------------------------------------------------- Office -> PDF
def office_to_pdf(src: Path, workdir: Path) -> list[Path]:
    produced = convert_with_soffice(src, output_dir(workdir), "pdf")
    return [produced]


# ------------------------------------------------------------ PDF -> Markdown
def pdf_to_markdown(src: Path, workdir: Path) -> list[Path]:
    """Reuses the same structure-aware HTML pipeline as the document editor
    (pdf2docx + LibreOffice when available, positioned-text fallback
    otherwise) and converts that HTML into Markdown."""
    from .html_edit import extract_html, html_to_markdown

    html = extract_html(src, workdir)
    markdown = html_to_markdown(html)
    out = output_dir(workdir) / f"{src.stem}.md"
    out.write_text(markdown, encoding="utf-8")
    return [out]


# ------------------------------------------------------------------ OCR PDF
def ocr_pdf(src: Path, workdir: Path, language: str = "eng") -> list[Path]:
    """Adds an invisible, selectable/searchable text layer over each page
    using PyMuPDF's built-in Tesseract integration. Requires the `tesseract`
    binary (and the requested language's tessdata) on the host; missing
    binary raises the same "unavailable on this deployment" 424 as other
    optional tools."""
    import fitz

    if not shutil_which("tesseract"):
        # Surfaced by the dispatcher as a 424. tesseract is a system binary, not
        # a pip package, so it must be present in the deployment image.
        raise ImportError(
            "the tesseract OCR engine is not installed on this server — the deployment "
            "image needs the 'tesseract-ocr' system package (plus language data such as "
            "'tesseract-ocr-kor' / 'tesseract-ocr-chi-sim' for Korean/Chinese)"
        )

    # PyMuPDF's pdfocr_* needs to know where tessdata lives. It reads
    # TESSDATA_PREFIX when tessdata=None, but apt/Homebrew installs don't always
    # export it — resolve the directory ourselves so OCR works out of the box
    # once the binary is present.
    tessdata = _resolve_tessdata()

    doc = _open_fitz(src)
    out = output_dir(workdir) / f"{src.stem}_ocr.pdf"
    try:
        ocr_doc = fitz.open()
        for page in doc:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            temp_pdf_bytes = pixmap.pdfocr_tobytes(language=language, tessdata=tessdata)
            page_doc = fitz.open("pdf", temp_pdf_bytes)
            ocr_doc.insert_pdf(page_doc)
            page_doc.close()
        ocr_doc.save(str(out), garbage=4, deflate=True)
        ocr_doc.close()
    except RuntimeError as exc:
        raise HTTPException(
            status_code=422,
            detail=f"OCR failed — check the language code '{language}' has tessdata installed ({exc}).",
        )
    finally:
        doc.close()
    return [out]


def _resolve_tessdata() -> str | None:
    """Locate the Tesseract `tessdata` directory (traineddata files).

    Returns None when it can't be found, in which case PyMuPDF falls back to its
    own TESSDATA_PREFIX handling. Order: explicit env var, then the standard
    apt/Homebrew locations, then a path derived from the tesseract binary."""
    import os

    env = os.environ.get("TESSDATA_PREFIX")
    if env and Path(env).is_dir():
        return env

    for candidate in (
        "/usr/share/tesseract-ocr/5/tessdata",
        "/usr/share/tesseract-ocr/4.00/tessdata",
        "/usr/share/tesseract-ocr/tessdata",
        "/usr/share/tessdata",
        "/usr/local/share/tessdata",
        "/opt/homebrew/share/tessdata",
    ):
        if Path(candidate).is_dir():
            return candidate

    binary = shutil_which("tesseract")
    if binary:
        share = Path(binary).resolve().parent.parent / "share"
        globbed = sorted(share.glob("tesseract-ocr/*/tessdata"))
        for path in (share / "tessdata", *globbed):
            if path.is_dir():
                return str(path)
    return None


def shutil_which(name: str) -> str | None:
    import shutil

    return shutil.which(name)


def image_to_text(inputs: list[Path], workdir: Path, fmt: str = "txt", language: str = "eng") -> list[Path]:
    """OCR one or more images (JPEG/PNG/HEIC/…) and emit a single document.

    Output differs by format, on purpose:
      * "txt"  — the recognised plain text only.
      * "pdf"  — a searchable scan: each page shows the ORIGINAL IMAGE with an
                 invisible, selectable text layer over it (PyMuPDF `pdfocr`).

    (A "docx" output was dropped — Word can't faithfully reproduce the image's
    text layout, so the results were unreliable.)

    HEIC/HEIF needs the optional `pillow-heif` decoder; without it those inputs
    raise a clear 422 while JPEG/PNG keep working."""
    import io

    import fitz
    from PIL import Image, ImageOps

    try:
        import pillow_heif  # optional native decoder for HEIC/HEIF

        pillow_heif.register_heif_opener()
    except Exception:  # noqa: BLE001 — HEIC simply stays unsupported without it
        pass

    fmt = (fmt or "txt").lower()
    if fmt not in {"txt", "pdf"}:
        raise HTTPException(status_code=422, detail="Choose an output format: TXT or PDF.")
    if not inputs:
        raise HTTPException(status_code=422, detail="Upload at least one image.")

    if not shutil_which("tesseract"):
        raise ImportError(
            "the tesseract OCR engine is not installed on this server — the deployment "
            "image needs the 'tesseract-ocr' system package (plus language data such as "
            "'tesseract-ocr-kor' / 'tesseract-ocr-chi-sim' for Korean/Chinese)"
        )
    tessdata = _resolve_tessdata()

    # Normalise every input to PNG bytes once (handles EXIF rotation + HEIC).
    images_png: list[bytes] = []
    for path in inputs:
        try:
            with Image.open(path) as im:
                im = ImageOps.exif_transpose(im) or im
                buf = io.BytesIO()
                im.convert("RGB").save(buf, "PNG")
                images_png.append(buf.getvalue())
        except Exception as exc:  # noqa: BLE001
            hint = ""
            if path.suffix.lower() in {".heic", ".heif"}:
                hint = " (HEIC support needs the 'pillow-heif' package on the server)"
            raise HTTPException(status_code=422, detail=f"'{path.name}' is not a readable image{hint} ({exc}).")

    stem = inputs[0].stem if len(inputs) == 1 else "images"
    out_dir = output_dir(workdir)

    # For PDF, keep the OCR page (image + text layer) as-is. For txt, pull the
    # recognised plain text.
    ocr_pdf_pages: list[bytes] = []
    texts: list[str] = []
    try:
        for png in images_png:
            img_doc = fitz.open(stream=png, filetype="png")
            pixmap = img_doc[0].get_pixmap(matrix=fitz.Matrix(2, 2))
            page_pdf = pixmap.pdfocr_tobytes(language=language, tessdata=tessdata)
            img_doc.close()
            if fmt == "pdf":
                ocr_pdf_pages.append(page_pdf)
            else:
                ocr_doc = fitz.open("pdf", page_pdf)
                texts.append(ocr_doc[0].get_text().strip())
                ocr_doc.close()
    except RuntimeError as exc:
        raise HTTPException(
            status_code=422,
            detail=f"OCR failed — check the language code '{language}' has tessdata installed ({exc}).",
        )

    if fmt == "pdf":
        combined = fitz.open()
        for page_pdf in ocr_pdf_pages:
            page_doc = fitz.open("pdf", page_pdf)
            combined.insert_pdf(page_doc)
            page_doc.close()
        out = out_dir / f"{stem}_text.pdf"
        combined.save(str(out), garbage=4, deflate=True)
        combined.close()
        return [out]

    out = out_dir / f"{stem}_text.txt"
    out.write_text("\n\n".join(t for t in texts if t), encoding="utf-8")
    return [out]


# -------------------------------------------------------------- Images -> PDF
def images_to_pdf(inputs: list[Path], workdir: Path, enhance: bool = False) -> list[Path]:
    """`enhance=True` (Scan to PDF) applies a light auto-contrast/grayscale
    pass so phone-camera photos of documents read like a real scan.

    Uses PyMuPDF (``fitz``) to build the PDF rather than img2pdf: img2pdf is not
    installed on the hosted platform, so JPG/Image-to-PDF 424'd with "No module
    named 'img2pdf'". fitz opens each image and embeds it losslessly, one page
    per image sized to that image."""
    import fitz
    from PIL import Image, ImageOps

    prepared: list[str] = []
    flat_dir = workdir / "flat"
    flat_dir.mkdir(exist_ok=True)
    for path in inputs:
        try:
            with Image.open(path) as im:
                if enhance:
                    im = ImageOps.exif_transpose(im) or im
                    im = ImageOps.autocontrast(im.convert("L"), cutoff=1).convert("RGB")
                    flat = flat_dir / f"{path.stem}.jpg"
                    im.save(flat, "JPEG", quality=92)
                    prepared.append(str(flat))
                # Flatten alpha/palette images onto a white background so
                # transparency doesn't render as black in the PDF.
                elif im.mode in ("RGBA", "LA", "P"):
                    rgba = im.convert("RGBA")
                    background = Image.new("RGB", rgba.size, (255, 255, 255))
                    background.paste(rgba, mask=rgba.split()[-1])
                    flat = flat_dir / f"{path.stem}.jpg"
                    background.save(flat, "JPEG", quality=92)
                    prepared.append(str(flat))
                else:
                    prepared.append(str(path))
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status_code=422, detail=f"'{path.name}' is not a readable image ({exc}).")

    out = output_dir(workdir) / ("scan.pdf" if enhance else "images.pdf")
    doc = fitz.open()
    try:
        for img_path in prepared:
            try:
                img_doc = fitz.open(img_path)
                pdf_bytes = img_doc.convert_to_pdf()
                img_doc.close()
            except Exception as exc:  # noqa: BLE001
                raise HTTPException(
                    status_code=422,
                    detail=f"'{Path(img_path).name}' could not be added to the PDF ({exc}).",
                )
            img_pdf = fitz.open("pdf", pdf_bytes)
            try:
                doc.insert_pdf(img_pdf)
            finally:
                img_pdf.close()
        doc.save(str(out))
    finally:
        doc.close()
    return [out]


# ------------------------------------------------------------- Image -> Image
_IMAGE_FORMATS = {"jpg": "JPEG", "jpeg": "JPEG", "png": "PNG", "webp": "WEBP", "bmp": "BMP", "tiff": "TIFF", "gif": "GIF"}


def convert_images(inputs: list[Path], workdir: Path, target: str) -> list[Path]:
    from PIL import Image

    target = target.lower().lstrip(".")
    if target not in _IMAGE_FORMATS:
        raise HTTPException(
            status_code=422,
            detail=f"'{target}' isn't a format this tool can create. Choose one of: {', '.join(sorted(set(_IMAGE_FORMATS)))}.",
        )
    fmt = _IMAGE_FORMATS[target]
    out = output_dir(workdir)
    outputs: list[Path] = []
    for path in inputs:
        try:
            source_ext = path.suffix.lower().lstrip(".")
            dest = out / f"{path.stem}.{target}"
            if _IMAGE_FORMATS.get(source_ext) == fmt:
                # Already in the requested format — return it unchanged instead
                # of a lossy re-encode round-trip through Pillow.
                shutil.copyfile(path, dest)
                outputs.append(dest)
                continue
            with Image.open(path) as im:
                if fmt == "JPEG" and im.mode in ("RGBA", "LA", "P"):
                    rgba = im.convert("RGBA")
                    background = Image.new("RGB", rgba.size, (255, 255, 255))
                    background.paste(rgba, mask=rgba.split()[-1])
                    im = background
                im.save(dest, fmt)
                outputs.append(dest)
        except HTTPException:
            raise
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status_code=422, detail=f"'{path.name}' could not be converted ({exc}).")
    return outputs


# ---------------------------------------------------------------- Text -> PDF
def text_to_pdf(src: Path, workdir: Path) -> list[Path]:
    import fitz

    try:
        content = src.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Could not read '{src.name}' ({exc}).")

    doc = fitz.open()
    rect = fitz.paper_rect("a4")
    margin = 54
    text_rect = fitz.Rect(rect.x0 + margin, rect.y0 + margin, rect.x1 - margin, rect.y1 - margin)
    remaining = content or " "
    while remaining:
        page = doc.new_page(width=rect.width, height=rect.height)
        overflow = page.insert_textbox(text_rect, remaining, fontsize=11, fontname="cour")
        # insert_textbox returns unused height (>=0 done) or negative when text
        # overflowed; fitz has no built-in pagination, so chunk by lines.
        if overflow >= 0:
            break
        lines = remaining.splitlines()
        # Estimate lines per page from font size and box height.
        per_page = max(1, int(text_rect.height / (11 * 1.2)) - 1)
        remaining = "\n".join(lines[per_page:])
        page_text = "\n".join(lines[:per_page])
        page.add_redact_annot(text_rect)  # clear the overflowed draw
        page.apply_redactions()
        page.insert_textbox(text_rect, page_text, fontsize=11, fontname="cour")

    out = output_dir(workdir) / f"{src.stem}.pdf"
    doc.save(str(out))
    doc.close()
    return [out]
