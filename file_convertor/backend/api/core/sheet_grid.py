"""Spreadsheets as a real grid: workbook -> HTML table -> workbook.

The document editor opens every other format through LibreOffice's PDF
render (see html_edit.extract_html), which is right for paginated documents
but wrong for a workbook: a rendered .xlsx is a flat picture of whatever
happened to fit Calc's print area, with no cells, no column letters and no
row numbers -- nothing a spreadsheet user can edit.

So spreadsheets take this path instead: openpyxl reads the real cell grid
(values, number formats, fonts, fills, alignment, borders, merges, column
widths) and it is emitted as one `<table class="xl-grid">` per worksheet,
with an A/B/C column header and 1/2/3 row gutter. Export reverses it --
the edited table is written back as a real workbook, not re-derived from a
PDF -- so a round trip keeps cells as cells.
"""

from __future__ import annotations

import csv
import datetime as dt
import html as html_lib
import logging
import re
from pathlib import Path
from typing import NamedTuple

logger = logging.getLogger(__name__)

SPREADSHEET_EXTS = {".xlsx", ".xlsm", ".xls", ".ods", ".csv"}
# Formats openpyxl cannot read: LibreOffice converts them to .xlsx first.
_NEEDS_CONVERSION = {".xls", ".ods"}

# A workbook can address a million rows; the editor is a DOM, not Excel.
# Bound what one sheet contributes so a sparse-but-huge sheet can't produce
# a payload the browser chokes on. Truncation is reported to the UI
# (`data-truncated`) rather than silently dropping content.
MAX_ROWS = 1000
MAX_COLS = 64
MAX_CELLS = 30_000
# A small sheet still renders as a grid, not as three lonely boxes.
MIN_ROWS = 24
MIN_COLS = 12

_DEFAULT_COL_PX = 80
_DEFAULT_ROW_PX = 22
_CHAR_TO_PX = 7  # Excel column width is in "0" characters, ~7px at 11pt Calibri


# --------------------------------------------------------------- reading

def column_letter(index: int) -> str:
    """1 -> A, 27 -> AA (openpyxl's get_column_letter, without the import)."""
    letters = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _as_xlsx(src: Path, workdir: Path) -> Path | None:
    """`src` in a form openpyxl can open, or None if it can't be produced."""
    ext = src.suffix.lower()
    if ext not in _NEEDS_CONVERSION:
        return src
    from .office import convert_with_soffice, find_soffice

    if not find_soffice():
        return None
    out = workdir / "sheet_xlsx"
    out.mkdir(exist_ok=True)
    try:
        return convert_with_soffice(src, out, "xlsx")
    except Exception:  # noqa: BLE001 — caller falls back to the generic path
        logger.info("Could not convert %s to xlsx for grid editing", src.name)
        return None


# `[$₹-en-IN]`, `"$"#,##0.00` and a bare `€` all name a currency; the quoted
# form is the common one, so the format body is unquoted before this runs.
_CURRENCY_RE = re.compile(r"\[\$([^\]\-]*)[^\]]*\]|([$€£¥₹₩])")
_DECIMALS_RE = re.compile(r"\.([0#]+)")
_DATE_TOKEN_RE = re.compile(r"(yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|AM/PM|am/pm)")
_PY_DATE_TOKEN = {
    "yyyy": "%Y", "yy": "%y",
    "mmmm": "%B", "mmm": "%b", "mm": "%m", "m": "%m",
    "dddd": "%A", "ddd": "%a", "dd": "%d", "d": "%d",
    "hh": "%H", "h": "%H", "ss": "%S", "s": "%S",
    "AM/PM": "%p", "am/pm": "%p",
}


def _format_datetime(value, fmt: str) -> str:
    """Render a date/time cell the way its number format asks. Only the common
    tokens are translated; anything else falls back to ISO, which is never
    wrong, just not what the author picked."""
    body = fmt.split(";")[0].replace('"', "").replace("\\", "")
    # `mm` means minutes, not months, when it directly follows h/hh or
    # precedes ss -- the one ambiguity in Excel's format language that
    # changes the meaning of a whole column.
    body = re.sub(r"(h{1,2}\]?\s*:\s*)mm", lambda m: m.group(1) + "%M", body, flags=re.IGNORECASE)
    body = re.sub(r"mm(\s*:\s*s{1,2})", lambda m: "%M" + m.group(1), body, flags=re.IGNORECASE)
    out = _DATE_TOKEN_RE.sub(lambda m: _PY_DATE_TOKEN[m.group(0)], body)
    if "%" not in out:
        return value.isoformat(sep=" ") if isinstance(value, dt.datetime) else value.isoformat()
    try:
        return value.strftime(out).strip()
    except ValueError:
        return value.isoformat()


def _trim_float(value: float) -> str:
    text = f"{value:.10f}".rstrip("0").rstrip(".")
    return text or "0"


def _format_number(value: float, fmt: str) -> str:
    """Apply the parts of an Excel number format that change what the reader
    sees: percent, thousands separators, fixed decimals, currency symbol."""
    body = (fmt or "General").split(";")[0].strip()
    # Literal text is quoted in a number format (`"$"#,##0.00`); unquote it so
    # the symbol is visible to the currency/thousands checks below.
    body = body.replace('"', "").replace("\\", "")
    if body in ("", "General", "@"):
        return _trim_float(value) if isinstance(value, float) else str(value)
    percent = "%" in body
    if percent:
        value = value * 100
    decimals_match = _DECIMALS_RE.search(body)
    decimals = len(decimals_match.group(1)) if decimals_match else 0
    thousands = "#,#" in body.replace(" ", "")
    text = f"{value:,.{decimals}f}" if thousands else f"{value:.{decimals}f}"
    currency = ""
    money = _CURRENCY_RE.search(body)
    if money:
        currency = money.group(1) or money.group(2) or ""
    if percent:
        text += "%"
    return f"{currency}{text}" if currency else text


def display_text(value, fmt: str = "General") -> str:
    """What Excel would show for `value` under number format `fmt`. Export
    calls this too, to tell an untouched cell from an edited one."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return _format_datetime(value, fmt)
    if isinstance(value, dt.timedelta):
        return str(value)
    if isinstance(value, (int, float)):
        return _format_number(value, fmt)
    return str(value)


def _display_value(cell) -> str:  # noqa: ANN001 — openpyxl Cell
    return display_text(cell.value, cell.number_format or "General")


class WorkbookStyle(NamedTuple):
    """Everything a cell's formatting resolves against: the Normal font every
    cell inherits, and the workbook's theme palette (an Excel-styled table
    stores its colors as theme references, not as RGB)."""

    font_name: str = "Calibri"
    font_size: float = 11.0
    font_color: "str | None" = None
    theme: tuple = ()


# Excel's `theme` index is NOT the order the theme file lists colors in: the
# first two pairs are swapped (index 0 is lt1, 1 is dk1, 2 is lt2, 3 is dk2).
_THEME_SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
                "accent4", "accent5", "accent6", "hlink", "folHlink")
_THEME_ORDER = (1, 0, 3, 2, 4, 5, 6, 7, 8, 9, 10, 11)
_DRAWING_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def theme_palette(wb) -> tuple:  # noqa: ANN001 — openpyxl Workbook
    """The workbook's theme colors as `#RRGGBB`, indexed the way a cell's
    `theme=` attribute indexes them. Empty when the workbook ships no theme."""
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return ()
    import xml.etree.ElementTree as ET

    try:
        scheme = ET.fromstring(raw).find(f"{_DRAWING_NS}themeElements/{_DRAWING_NS}clrScheme")
    except ET.ParseError:
        return ()
    if scheme is None:
        return ()
    found: dict[str, str] = {}
    for slot in scheme:
        name = slot.tag.rsplit("}", 1)[-1]
        srgb = slot.find(f"{_DRAWING_NS}srgbClr")
        sys_clr = slot.find(f"{_DRAWING_NS}sysClr")
        value = None
        if srgb is not None:
            value = srgb.get("val")
        elif sys_clr is not None:
            value = sys_clr.get("lastClr")
        if value and len(value) == 6:
            found[name] = "#" + value.upper()
    ordered = [found.get(_THEME_SLOTS[i], "") for i in _THEME_ORDER]
    return tuple(ordered)


def _apply_tint(hex_color: str, tint: float) -> str:
    """Excel stores a shade/highlight of a theme color as a tint on its
    luminance — without applying it, every banded row reads as the same flat
    color as its header."""
    if not tint:
        return hex_color
    import colorsys

    r, g, b = (int(hex_color[i : i + 2], 16) / 255 for i in (1, 3, 5))
    hue, lum, sat = colorsys.rgb_to_hls(r, g, b)
    lum = lum * (1 - tint) + tint if tint > 0 else lum * (1 + tint)
    r, g, b = colorsys.hls_to_rgb(hue, max(0.0, min(1.0, lum)), sat)
    return "#{:02X}{:02X}{:02X}".format(round(r * 255), round(g * 255), round(b * 255))


def _css_color(color, style: WorkbookStyle) -> str | None:  # noqa: ANN001 — openpyxl Color
    """`#RRGGBB` for a cell color of any flavour Excel writes: direct RGB, a
    theme reference (with its tint), or a legacy indexed-palette entry."""
    if color is None:
        return None
    kind = getattr(color, "type", None)
    if kind == "rgb":
        rgb = getattr(color, "rgb", None)
        if isinstance(rgb, str) and len(rgb) == 8 and rgb[:2] != "00":
            return _apply_tint("#" + rgb[2:], getattr(color, "tint", 0.0) or 0.0)
        return None
    if kind == "theme":
        index = getattr(color, "theme", None)
        if isinstance(index, int) and 0 <= index < len(style.theme) and style.theme[index]:
            return _apply_tint(style.theme[index], getattr(color, "tint", 0.0) or 0.0)
        return None
    if kind == "indexed":
        from openpyxl.styles.colors import COLOR_INDEX

        index = getattr(color, "indexed", None)
        # 64/65 are "automatic" foreground/background — no fixed color.
        if isinstance(index, int) and 0 <= index < len(COLOR_INDEX) and index not in (64, 65):
            value = COLOR_INDEX[index]
            if isinstance(value, str) and len(value) == 8:
                return "#" + value[2:]
        return None
    return None


_BORDER_WIDTH = {
    "hair": "1px", "thin": "1px", "medium": "2px", "thick": "3px",
    "dotted": "1px", "dashed": "1px", "double": "3px",
    "mediumDashed": "2px", "dashDot": "1px", "mediumDashDot": "2px",
    "dashDotDot": "1px", "mediumDashDotDot": "2px", "slantDashDot": "2px",
}
_BORDER_STYLE = {
    "dotted": "dotted", "dashed": "dashed", "double": "double",
    "mediumDashed": "dashed", "dashDot": "dashed", "mediumDashDot": "dashed",
    "dashDotDot": "dashed", "mediumDashDotDot": "dashed", "slantDashDot": "dashed",
}


def _border_css(side, style: WorkbookStyle) -> str | None:  # noqa: ANN001 — openpyxl Side
    line = getattr(side, "style", None)
    if not line:
        return None
    color = _css_color(getattr(side, "color", None), style) or "#9aa0a6"
    return f"{_BORDER_WIDTH.get(line, '1px')} {_BORDER_STYLE.get(line, 'solid')} {color}"


def workbook_style(wb) -> WorkbookStyle:  # noqa: ANN001 — openpyxl Workbook
    """The workbook's Normal font (inherited by every cell, so naming it on
    each `<td>` would only bloat the payload) plus its theme palette."""
    palette = theme_palette(wb)
    name, size, color = "Calibri", 11.0, None
    try:
        normal = wb._named_styles["Normal"].font
        name, size = normal.name or name, float(normal.sz or size)
        # Excel records the default text color explicitly (usually a theme
        # reference to black). Emitting it on every cell would put an inline
        # style on blank padding cells too, which export then has to save.
        color = _css_color(normal.color, WorkbookStyle(theme=palette))
    except Exception:  # noqa: BLE001 — openpyxl internals; the defaults are safe
        pass
    return WorkbookStyle(name, size, color, palette)


def _cell_css(cell, style: WorkbookStyle = WorkbookStyle()) -> str:  # noqa: ANN001
    """Inline style carrying this cell's real formatting into the editor."""
    parts: list[str] = []
    font = cell.font
    if font is not None:
        if font.bold:
            parts.append("font-weight:700")
        if font.italic:
            parts.append("font-style:italic")
        if font.underline and font.underline != "none":
            parts.append("text-decoration:underline")
        if font.strike:
            parts.append("text-decoration:line-through")
        if font.sz and float(font.sz) != style.font_size:
            parts.append(f"font-size:{float(font.sz):g}pt")
        if font.name and font.name != style.font_name:
            # Single quotes: this lands inside a double-quoted attribute.
            parts.append(f"font-family:'{font.name}'")
        color = _css_color(font.color, style)
        if color and color != style.font_color:
            parts.append(f"color:{color}")
    fill = cell.fill
    if fill is not None and getattr(fill, "patternType", None) == "solid":
        bg = _css_color(fill.fgColor, style)
        if bg:
            parts.append(f"background-color:{bg}")
    align = cell.alignment
    horizontal = getattr(align, "horizontal", None) if align else None
    if horizontal in ("left", "right", "center", "justify"):
        parts.append(f"text-align:{horizontal}")
    elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
        parts.append("text-align:right")  # Excel's "General" default for numbers
    # Only when the author set one: an unformatted cell must carry NO inline
    # style at all, so export can recognise the grid's blank padding cells and
    # leave them out of the saved workbook.
    vertical = getattr(align, "vertical", None) if align else None
    if vertical in ("top", "center", "bottom"):
        parts.append("vertical-align:" + {"center": "middle"}.get(vertical, vertical))
    if align is not None and align.wrap_text:
        parts.append("white-space:pre-wrap")
    border = cell.border
    if border is not None:
        for name, side in (("top", border.top), ("right", border.right), ("bottom", border.bottom), ("left", border.left)):
            css = _border_css(side, style)
            if css:
                parts.append(f"border-{name}:{css}")
    return ";".join(parts)


def _value_attrs(cell, formula: "str | None") -> str:  # noqa: ANN001 — openpyxl Cell
    """The cell's REAL value, alongside the text the editor shows. A date
    renders as "03 Jul 2026" and a price as "$19.99"; without the underlying
    value recorded here, saving the document would turn every one of them
    into a string. Export reads these back and only re-parses the text of
    cells whose displayed value actually changed."""
    if formula:
        # A formula's displayed text is its last cached result; keep the
        # formula itself so an untouched cell stays live on the way out.
        return (
            f' data-t="f" data-v="{html_lib.escape(formula, quote=True)}"'
            f' data-d="{html_lib.escape(_display_value(cell), quote=True)}"'
        )
    value = cell.value
    if isinstance(value, bool):
        tag, raw = "b", ("1" if value else "0")
    elif isinstance(value, (int, float)):
        tag, raw = "n", repr(value)
    elif isinstance(value, dt.datetime):
        tag, raw = "dt", value.isoformat()
    elif isinstance(value, dt.date):
        tag, raw = "d", value.isoformat()
    elif isinstance(value, dt.time):
        tag, raw = "tm", value.isoformat()
    else:
        return ""  # plain text (or empty): the cell's text IS its value
    fmt = cell.number_format or "General"
    attrs = f' data-t="{tag}" data-v="{html_lib.escape(raw, quote=True)}"'
    if fmt not in ("", "General"):
        attrs += f' data-fmt="{html_lib.escape(fmt, quote=True)}"'
    return attrs


def _used_bounds(ws) -> tuple[int, int]:  # noqa: ANN001 — openpyxl Worksheet
    """Real used size. `ws.max_row`/`max_column` count cells that only ever
    carried formatting, which for some exports is the entire sheet -- scan
    (bounded) for the last row/column holding an actual value."""
    max_row = min(ws.max_row or 1, MAX_ROWS * 4)
    max_col = min(ws.max_column or 1, MAX_COLS * 4)
    last_row = 0
    last_col = 0
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.value not in (None, ""):
                last_row = max(last_row, cell.row)
                last_col = max(last_col, cell.column)
    for rng in ws.merged_cells.ranges:  # a merge defines used space on its own
        last_row = max(last_row, rng.max_row)
        last_col = max(last_col, rng.max_col)
    # The scan stopped at its own bound, so there is more sheet beyond it:
    # report the sheet's declared size instead, or the "showing N of M rows"
    # notice would quote the scan limit as the sheet's true size.
    if last_row >= max_row:
        last_row = max(last_row, ws.max_row or last_row)
    if last_col >= max_col:
        last_col = max(last_col, ws.max_column or last_col)
    return last_row, last_col


def _sheet_html(ws, name: str, formulas: dict | None = None, style: WorkbookStyle = WorkbookStyle()) -> str:  # noqa: ANN001
    """One worksheet as an editable HTML grid. `formulas` maps (row, col) ->
    formula text for the cells that hold one."""
    used_rows, used_cols = _used_bounds(ws)
    rows = max(min(used_rows, MAX_ROWS), MIN_ROWS)
    cols = max(min(used_cols, MAX_COLS), MIN_COLS)
    if rows * cols > MAX_CELLS:  # keep the payload sane, rows first
        rows = max(MIN_ROWS, MAX_CELLS // cols)
    truncated = used_rows > rows or used_cols > cols

    # Merged ranges: the top-left cell spans the block, the rest are dropped.
    spans: dict[tuple[int, int], tuple[int, int]] = {}
    covered: set[tuple[int, int]] = set()
    for rng in ws.merged_cells.ranges:
        spans[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))

    out: list[str] = []
    gridlines = getattr(ws.sheet_view, "showGridLines", True) is not False
    out.append(
        f'<div class="xl-sheet" data-sheet-name="{html_lib.escape(name, quote=True)}"'
        f' data-rows="{rows}" data-cols="{cols}"'
        f' data-truncated="{"true" if truncated else "false"}"'
        f' data-used-rows="{used_rows}" data-used-cols="{used_cols}">'
    )
    out.append(f'<table class="xl-grid{"" if gridlines else " xl-no-lines"}">')

    # <colgroup> carries the real column widths, so the sheet looks like the
    # workbook rather than auto-sizing to its content.
    out.append('<colgroup><col class="xl-gutter-col">')
    for c in range(1, cols + 1):
        dim = ws.column_dimensions.get(column_letter(c))
        width = getattr(dim, "width", None) if dim is not None else None
        px = round(width * _CHAR_TO_PX) if width else _DEFAULT_COL_PX
        out.append(f'<col style="width:{max(24, min(480, px))}px">')
    out.append("</colgroup>")

    out.append('<thead><tr><th class="xl-corner" contenteditable="false"></th>')
    for c in range(1, cols + 1):
        out.append(f'<th class="xl-colhead" contenteditable="false">{column_letter(c)}</th>')
    out.append("</tr></thead><tbody>")

    for r in range(1, rows + 1):
        dim = ws.row_dimensions.get(r)
        height = getattr(dim, "height", None) if dim is not None else None
        px = round(height * 96 / 72) if height else _DEFAULT_ROW_PX
        out.append(f'<tr style="height:{max(16, min(400, px))}px">')
        out.append(f'<th class="xl-rowhead" contenteditable="false">{r}</th>')
        for c in range(1, cols + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            formula = (formulas or {}).get((r, c))
            attrs = f' data-cell="{column_letter(c)}{r}"' + _value_attrs(cell, formula)
            span = spans.get((r, c))
            if span:
                rowspan, colspan = span
                if rowspan > 1:
                    attrs += f' rowspan="{rowspan}"'
                if colspan > 1:
                    attrs += f' colspan="{colspan}"'
            css = _cell_css(cell, style)
            if css:
                attrs += f' style="{html_lib.escape(css, quote=True)}"'
            text = _display_value(cell) or (formula or "")
            out.append(f"<td{attrs}>{html_lib.escape(text)}</td>")
        out.append("</tr>")
    out.append("</tbody></table></div>")
    return "".join(out)


def _csv_workbook(src: Path):
    """A .csv as a one-sheet in-memory workbook, so it takes the same path."""
    from openpyxl import Workbook

    text = src.read_text(encoding="utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    wb = Workbook()
    ws = wb.active
    ws.title = src.stem[:31] or "Sheet1"
    for row in csv.reader(text.splitlines(), dialect):
        ws.append([c if c != "" else None for c in row[:MAX_COLS]])
        if ws.max_row >= MAX_ROWS:
            break
    return wb


def _formula_map(path: Path) -> dict[str, dict[tuple[int, int], str]]:
    """{sheet name: {(row, col): formula}} — read from a second, non-cached
    open of the workbook, because `data_only=True` (which gives the results
    the author last saw) replaces every formula with its cached value."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=False)
    except Exception:  # noqa: BLE001 — cached values alone still render fine
        return {}
    out: dict[str, dict[tuple[int, int], str]] = {}
    for ws in wb.worksheets:
        found: dict[tuple[int, int], str] = {}
        for row in ws.iter_rows(
            min_row=1, max_row=min(ws.max_row or 1, MAX_ROWS),
            max_col=min(ws.max_column or 1, MAX_COLS),
        ):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    found[(cell.row, cell.column)] = cell.value
        if found:
            out[ws.title] = found
    wb.close()
    return out


def spreadsheet_to_html(src: Path, workdir: Path) -> str | None:
    """Every worksheet in `src` as an editable HTML grid, or None when the
    file can't be read as a workbook (caller falls back to the generic
    document pipeline)."""
    ext = src.suffix.lower()
    try:
        if ext == ".csv":
            wb = _csv_workbook(src)
            formulas: dict[str, dict[tuple[int, int], str]] = {}
        else:
            path = _as_xlsx(src, workdir)
            if path is None:
                return None
            from openpyxl import load_workbook

            # `data_only=True` gives each formula's last calculated result —
            # what the author actually saw in Excel.
            wb = load_workbook(path, data_only=True)
            formulas = _formula_map(path)
    except Exception as exc:  # noqa: BLE001 — not a readable workbook
        logger.info("Spreadsheet grid extraction failed for %s: %s", src.name, exc)
        return None

    sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"] or list(wb.worksheets)
    if not sheets:
        return None
    style = workbook_style(wb)
    return "".join(
        _sheet_html(ws, ws.title, formulas.get(ws.title), style) for ws in sheets
    )


# --------------------------------------------------------------- writing

_NUMERIC_RE = re.compile(r"^-?[\d,]*\.?\d+%?$")
_CURRENCY_CHARS = "$€£¥₹₩"


def _parse_cell_text(text: str) -> tuple:
    """Text typed into the grid, back to (value, number_format) — so a number
    stays a number (sortable, summable, chartable) instead of becoming a
    string, and the currency/percent the user typed survives as real Excel
    formatting rather than as part of the text."""
    stripped = text.strip()
    if stripped == "":
        return None, None
    if stripped.startswith("="):
        return stripped, None  # a formula the user typed, stored as one
    currency = ""
    body = stripped
    if body[:1] in _CURRENCY_CHARS:
        currency, body = body[0], body[1:].strip()
    elif body[-1:] in _CURRENCY_CHARS:
        currency, body = body[-1], body[:-1].strip()
    if not _NUMERIC_RE.match(body):
        return text, None
    percent = body.endswith("%")
    digits = body.rstrip("%").replace(",", "")
    try:
        number = float(digits)
    except ValueError:
        return text, None
    decimals = len(digits.partition(".")[2])
    thousands = "," in body
    if percent:
        number = number / 100
        return number, f"0.{'0' * decimals}%" if decimals else "0%"
    if number.is_integer() and not decimals:
        number = int(number)
    integer_part = "#,##0" if thousands else "0"
    if currency:
        return number, f'"{currency}"{integer_part}{"." + "0" * decimals if decimals else ""}'
    if thousands or decimals:
        return number, f"{integer_part}{'.' + '0' * decimals if decimals else ''}"
    return number, None


def _revive(tag: str, raw: str):
    """The original typed value an editor cell was built from."""
    try:
        if tag == "n":
            return int(raw) if re.fullmatch(r"-?\d+", raw) else float(raw)
        if tag == "b":
            return raw == "1"
        if tag == "dt":
            return dt.datetime.fromisoformat(raw)
        if tag == "d":
            return dt.date.fromisoformat(raw)
        if tag == "tm":
            return dt.time.fromisoformat(raw)
    except ValueError:
        return None
    return None


def _cell_value(td, text: str) -> tuple:
    """(value, number_format) for one edited grid cell. An untouched cell —
    its text still matches what its recorded value renders as — is written
    back with that exact value and format, so a date stays a date and
    `=D2*E2` stays a live formula. Only genuinely edited text is re-parsed."""
    tag = td.get("data-t") or ""
    raw = td.get("data-v")
    fmt = td.get("data-fmt")
    if tag == "f" and raw is not None:
        if text.strip() == (td.get("data-d") or "").strip() or text.strip() == raw.strip():
            return raw, None  # untouched (or the formula itself was shown)
        return _parse_cell_text(text)
    if tag and raw is not None:
        original = _revive(tag, raw)
        if original is not None and display_text(original, fmt or "General").strip() == text.strip():
            return original, fmt
        value, parsed_fmt = _parse_cell_text(text)
        # An edited number keeps the column's own format ("$19.99" -> "$25.00"),
        # since the symbol the user sees comes from the format, not the text.
        if fmt and isinstance(value, (int, float)) and not isinstance(value, bool):
            return value, fmt
        return value, parsed_fmt
    return _parse_cell_text(text)


_RGB_FUNC_RE = re.compile(r"rgba?\(\s*([\d.]+)\s*,\s*([\d.]+)\s*,\s*([\d.]+)(?:\s*,\s*([\d.]+))?\s*\)")


def _css_to_argb(value: str) -> "str | None":
    """A CSS colour from the editor as Excel's `AARRGGBB`. Both spellings have
    to be understood: the grid ships `#RRGGBB`, but the moment a cell's style
    is touched through the DOM (any toolbox colour, bold, alignment...) the
    browser reserialises the whole declaration into `rgb(r, g, b)` form."""
    text = (value or "").strip().lower()
    if not text or text in ("transparent", "inherit", "initial", "none"):
        return None
    if text.startswith("#"):
        digits = text[1:]
        if len(digits) == 3:
            digits = "".join(ch * 2 for ch in digits)
        if len(digits) == 6:
            return "FF" + digits.upper()
        return None
    match = _RGB_FUNC_RE.match(text)
    if not match:
        return None
    if match.group(4) is not None and float(match.group(4)) == 0:
        return None  # fully transparent: no fill at all
    try:
        r, g, b = (min(255, max(0, round(float(match.group(i))))) for i in (1, 2, 3))
    except ValueError:
        return None
    return f"FF{r:02X}{g:02X}{b:02X}"


# CSS border width -> the nearest Excel line weight, and the dash styles that
# have an Excel equivalent.
_CSS_BORDER_RE = re.compile(r"([\d.]+)px\s+(\w+)(?:\s+(.*))?")
_CSS_BORDER_STYLE = {"dashed": "dashed", "dotted": "dotted", "double": "double"}


def _css_side(value: str):
    """One `border-<side>` declaration as an openpyxl Side, or None."""
    from openpyxl.styles import Side

    match = _CSS_BORDER_RE.match((value or "").strip())
    if not match:
        return None
    width, line, color = float(match.group(1)), match.group(2).lower(), match.group(3) or ""
    if line in ("none", "hidden"):
        return None
    style = _CSS_BORDER_STYLE.get(line)
    if style is None:
        style = "thin" if width <= 1 else ("medium" if width <= 2 else "thick")
    argb = _css_to_argb(color)
    return Side(style=style, color=argb) if argb else Side(style=style)


def _style_lookup(style: str) -> dict:
    out: dict[str, str] = {}
    for part in (style or "").split(";"):
        if ":" in part:
            key, _, value = part.partition(":")
            out[key.strip().lower()] = value.strip()
    return out


def grid_pages_to_workbook(pages: list[dict], out_path: Path) -> Path:
    """Edited grid pages (the editor's `fc-page` wrappers, kind="sheet") back
    into a real .xlsx — values, bold/italic, alignment, fills and merges.
    Written directly rather than re-imported from a PDF render, so the result
    is a workbook with live cells."""
    from lxml import html as lxml_html
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill
    from openpyxl.utils import column_index_from_string

    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()

    for index, page in enumerate(pages):
        title = (page.get("name") or f"Sheet{index + 1}")[:31] or f"Sheet{index + 1}"
        while title in used_titles:  # Excel rejects duplicate sheet names
            title = f"{title[:28]}_{index + 1}"
        used_titles.add(title)
        ws = wb.create_sheet(title=title)
        try:
            root = lxml_html.fromstring("<div>" + page["inner"] + "</div>")
        except Exception:  # noqa: BLE001 — an unparseable page becomes an empty sheet
            continue

        fallback_row = 0
        for tr in root.iter("tr"):
            fallback_row += 1
            fallback_col = 0
            for td in tr:
                if td.tag != "td":
                    continue  # the A/B/C and 1/2/3 headers are <th>, never data
                fallback_col += 1
                ref = td.get("data-cell") or ""
                match = re.fullmatch(r"([A-Z]+)(\d+)", ref)
                if match:
                    col = column_index_from_string(match.group(1))
                    row = int(match.group(2))
                else:
                    col, row = fallback_col, fallback_row
                text = td.text_content().replace("\xa0", " ")
                style = _style_lookup(td.get("style") or "")
                rowspan = int(td.get("rowspan") or 1)
                colspan = int(td.get("colspan") or 1)
                value, number_format = _cell_value(td, text)
                if value is None and not style and rowspan == 1 and colspan == 1:
                    # The grid is padded out to a minimum size so it looks
                    # like a spreadsheet; writing those blanks would give the
                    # saved workbook a used range full of empty cells.
                    fallback_col += colspan - 1
                    continue
                cell = ws.cell(row=row, column=col)
                cell.value = value
                if number_format:
                    cell.number_format = number_format

                font_kwargs: dict = {}
                if style.get("font-weight") in ("700", "bold"):
                    font_kwargs["bold"] = True
                if style.get("font-style") == "italic":
                    font_kwargs["italic"] = True
                # Browsers may keep the shorthand or split it into the
                # longhand property; the cell is underlined either way.
                decoration = style.get("text-decoration") or style.get("text-decoration-line") or ""
                if "underline" in decoration:
                    font_kwargs["underline"] = "single"
                if "line-through" in decoration:
                    font_kwargs["strike"] = True
                color = _css_to_argb(style.get("color", ""))
                if color:
                    font_kwargs["color"] = color
                size = style.get("font-size", "")
                if size.endswith("pt"):
                    try:
                        font_kwargs["size"] = float(size[:-2])
                    except ValueError:
                        pass
                if font_kwargs:
                    cell.font = Font(**font_kwargs)
                background = _css_to_argb(style.get("background-color", ""))
                if background:
                    cell.fill = PatternFill(fill_type="solid", start_color=background, end_color=background)
                horizontal = style.get("text-align")
                wrap = style.get("white-space") == "pre-wrap"
                if horizontal in ("left", "right", "center", "justify") or wrap:
                    cell.alignment = Alignment(horizontal=horizontal, wrap_text=wrap or None)
                vertical = {"middle": "center", "top": "top", "bottom": "bottom"}.get(
                    style.get("vertical-align", "")
                )
                if vertical:
                    cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal,
                        vertical=vertical,
                        wrap_text=cell.alignment.wrap_text,
                    )

                # Ruled cells stay ruled: without this, every border the
                # workbook came with is lost the first time it is saved.
                sides = {
                    name: _css_side(style.get(f"border-{name}", "") or style.get("border", ""))
                    for name in ("top", "right", "bottom", "left")
                }
                if any(sides.values()):
                    cell.border = Border(**sides)

                if rowspan > 1 or colspan > 1:
                    ws.merge_cells(
                        start_row=row, start_column=col,
                        end_row=row + rowspan - 1, end_column=col + colspan - 1,
                    )
                fallback_col += colspan - 1

    if not wb.sheetnames:
        wb.create_sheet(title="Sheet1")
    wb.save(out_path)
    return out_path


def grid_page_to_csv(page: dict, out_path: Path) -> Path:
    """One edited grid page as .csv — the cells exactly as the editor shows
    them. Taken from the grid rather than from a saved workbook because a
    .csv holds text, not formulas or number formats, and the displayed text
    is what a .csv upload expects to get back."""
    from lxml import html as lxml_html

    rows: list[list[str]] = []
    try:
        root = lxml_html.fromstring("<div>" + page["inner"] + "</div>")
    except Exception:  # noqa: BLE001 — unparseable page: an empty csv
        root = None
    if root is not None:
        for tr in root.iter("tr"):
            row: list[str] = []
            cells = [td for td in tr if td.tag == "td"]
            if not cells:
                continue  # the A/B/C header row: not part of the data
            for td in cells:
                row.append(td.text_content().replace("\xa0", " ").strip())
                row.extend([""] * (int(td.get("colspan") or 1) - 1))
            while row and row[-1] == "":  # drop the grid's blank padding
                row.pop()
            rows.append(row)
    while rows and not rows[-1]:
        rows.pop()
    with out_path.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows(rows)
    return out_path
