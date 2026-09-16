"""Extract plain text from any supported document for the AI tools."""

from __future__ import annotations

import csv
from pathlib import Path

from fastapi import HTTPException

TEXT_EXTS = {".txt", ".md", ".log"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
PPTX_EXTS = {".pptx"}
XLSX_EXTS = {".xlsx"}
CSV_EXTS = {".csv"}
# Formats LibreOffice converts to text for us.
SOFFICE_EXTS = {".doc", ".odt", ".rtf", ".odp", ".ppt", ".ods", ".xls"}

SUPPORTED_EXTS = TEXT_EXTS | PDF_EXTS | DOCX_EXTS | PPTX_EXTS | XLSX_EXTS | CSV_EXTS | SOFFICE_EXTS


def extract_text(src: Path, workdir: Path) -> str:
    ext = src.suffix.lower()

    if ext in TEXT_EXTS:
        text = src.read_text(encoding="utf-8", errors="replace")
    elif ext in PDF_EXTS:
        text = _from_pdf(src)
    elif ext in DOCX_EXTS:
        text = _from_docx(src)
    elif ext in PPTX_EXTS:
        text = _from_pptx(src)
    elif ext in XLSX_EXTS:
        text = _from_xlsx(src)
    elif ext in CSV_EXTS:
        text = _from_csv(src)
    elif ext in SOFFICE_EXTS:
        text = _via_soffice(src, workdir)
    else:
        raise HTTPException(
            status_code=422,
            detail=f"'{src.name}' can't be used here — this tool works with: {', '.join(sorted(SUPPORTED_EXTS))}.",
        )

    text = text.strip()
    if not text:
        raise HTTPException(
            status_code=422,
            detail="No readable text found in this document (it may be a scanned image).",
        )
    return text


def _from_pdf(src: Path) -> str:
    import fitz

    try:
        doc = fitz.open(str(src))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"'{src.name}' is not a valid PDF ({exc}).")
    if doc.needs_pass:
        raise HTTPException(status_code=422, detail="PDF is password-protected. Unlock it first.")
    text = "\n\n".join(page.get_text() for page in doc)
    doc.close()
    return text


def _from_docx(src: Path) -> str:
    from docx import Document

    try:
        document = Document(str(src))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Could not read '{src.name}' ({exc}).")
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _from_pptx(src: Path) -> str:
    from pptx import Presentation

    try:
        prs = Presentation(str(src))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Could not read '{src.name}' ({exc}).")
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _from_xlsx(src: Path) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(src), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Could not read '{src.name}' ({exc}).")
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)


def _from_csv(src: Path) -> str:
    try:
        with src.open(newline="", encoding="utf-8", errors="replace") as fh:
            return "\n".join("\t".join(row) for row in csv.reader(fh))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Could not read '{src.name}' ({exc}).")


def _via_soffice(src: Path, workdir: Path) -> str:
    from .office import convert_with_soffice

    out_dir = workdir / "txt"
    out_dir.mkdir(exist_ok=True)
    produced = convert_with_soffice(src, out_dir, "txt:Text (encoded):UTF8")
    return produced.read_text(encoding="utf-8", errors="replace")
