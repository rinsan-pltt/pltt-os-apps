"""AI document tools: summarize, proofread/correct, translate, and export.

POST /ai/summarize  — file + length option → JSON {summary, model, ...}
POST /ai/proofread  — file → JSON {corrected_text, fixes[], model, ...}
POST /ai/translate  — file + target_language → JSON {html, model, ...}
POST /ai/export     — text + format (pdf|txt|docx) → file download
POST /ai/assistant  — the editor's outline + a prompt → JSON {answer, operations}
POST /ai/chat/open  — file → JSON {text, filename, ...}: read once for the document chat
POST /ai/chat       — the document's text + conversation + prompt → JSON {answer}
POST /ai/agent      — the global chat: tool catalog + files + prompt → JSON {reply, steps}
"""

from __future__ import annotations

import json
import re
from typing import Any

import anyio
from fastapi import APIRouter, Form, HTTPException, UploadFile

from pydantic import BaseModel, Field

from ..core import assistant
from ..core import chat_agent
from ..core import proofread as proofreading
from ..core.files import cleanup, new_workdir, respond_with, safe_name, save_upload
from ..core.html_edit import (
    EDIT_SUPPORTED_EXTS,
    EXTRACTION_TIMEOUT_S,
    extract_data_images,
    extract_html_bounded,
    restore_data_images,
)
from ..core.llm import DEFAULT_MODEL, call_openai
from ..core.palette import require_permission
from ..core.secrets import read_secret
from ..core.text_extract import SUPPORTED_EXTS, extract_text

try:
    from fastapi import Depends
    from palette_sdk import get_plugin_context  # type: ignore

    _CTX_DEP: Any = Depends(get_plugin_context)
except ImportError:  # standalone dev/tests — secrets come from os.environ
    _CTX_DEP = None

router = APIRouter(tags=["ai"])

SUMMARIZE_MAX_CHARS = 120_000
# The whole document is always translated (no truncation). It's split into
# element-boundary chunks of roughly this many characters so each LLM call stays
# well within output limits, and the chunks are translated concurrently.
TRANSLATE_CHUNK_CHARS = 16_000
TRANSLATE_MAX_CONCURRENCY = 4


def _require_openai_key(ctx: Any) -> str:
    """Fail fast, before spending time on extraction, if there's no key."""
    api_key = read_secret(ctx, "OPENAI_KEY")
    if not api_key:
        raise HTTPException(
            status_code=424,
            detail=(
                "OPENAI_KEY is not configured. Add it to the plugin's .env "
                "(local) or plugin secrets (hosted) and restart."
            ),
        )
    return api_key

_CODE_FENCE_RE = re.compile(r"^```(?:html)?\s*(.*?)\s*```$", re.DOTALL)


def _strip_code_fences(text: str) -> str:
    text = text.strip()
    match = _CODE_FENCE_RE.match(text)
    return match.group(1) if match else text


# --------------------------------------------------------------- Translation

def _translate_system_prompt(language: str) -> str:
    return (
        f"You translate HTML documents into {language}. Translate ONLY the human-readable "
        "text content and attributes like alt/title/placeholder. NEVER change, add, remove, "
        "merge, split or reorder any tags — copy the tag structure byte-for-byte, including "
        "empty elements such as <td></td>, <th></th> and <tr></tr>. Preserve every table, row "
        "and cell exactly and keep empty cells empty. Do not change tag names, style/class/id "
        "attributes, or any __IMG_PLACEHOLDER_n__ token. You are given one fragment of a larger "
        "document — translate it as-is and respond with only the translated HTML fragment, no "
        "commentary and no code fences."
    )


def _split_body(html: str) -> tuple[str, str, str]:
    """Split into (prefix, body_inner, suffix). The prefix (doctype/head/style)
    and suffix are passed through untranslated so CSS/table styling survives."""
    low = html.lower()
    i = low.find("<body")
    if i == -1:
        return "", html, ""
    j = html.find(">", i)
    k = low.rfind("</body>")
    if j == -1 or k == -1 or k < j:
        return "", html, ""
    return html[: j + 1], html[j + 1 : k], html[k:]


def _shallow_tags(el) -> tuple[str, str]:
    """Opening and closing tag of `el` with no content, e.g. ('<table ...>', '</table>')."""
    import copy

    from lxml import etree

    clone = copy.deepcopy(el)
    clone.tail = None
    clone.text = None
    for child in list(clone):
        clone.remove(child)
    serialized = etree.tostring(clone, encoding="unicode", method="html")
    idx = serialized.rfind("</")
    if idx == -1:  # void/self-closing element — can't wrap around its children
        return "", ""
    return serialized[:idx], serialized[idx:]


def _chunk_container(container, limit: int) -> list[str]:
    """Serialise a container's children into HTML chunks each ≤ `limit` chars,
    cutting only BETWEEN whole elements. An element larger than the limit is
    split by recursing into its own children and wrapping each sub-chunk back in
    a shallow copy of that element — so a big table splits into valid smaller
    tables (rows and empty cells intact), never mid-tag."""
    from lxml import etree

    chunks: list[str] = []
    cur = container.text or ""
    for child in container:
        piece = etree.tostring(child, encoding="unicode", method="html")  # includes child.tail
        if len(piece) > limit and len(child) > 0:
            open_tag, close_tag = _shallow_tags(child)
            if open_tag:
                if cur.strip():
                    chunks.append(cur)
                cur = ""
                sub_limit = max(2000, limit - len(open_tag) - len(close_tag))
                for sub in _chunk_container(child, sub_limit):
                    chunks.append(open_tag + sub + close_tag)
                cur = child.tail or ""
                continue
        if cur and len(cur) + len(piece) > limit:
            chunks.append(cur)
            cur = ""
        cur += piece
    if cur.strip():
        chunks.append(cur)
    return chunks


def _split_html_chunks(html: str, limit: int) -> tuple[str, list[str], str]:
    prefix, inner, suffix = _split_body(html)
    if len(inner) <= limit:
        return prefix, ([inner] if inner.strip() else []), suffix
    from lxml import html as lxml_html

    try:
        wrapper = lxml_html.fromstring("<div>" + inner + "</div>")
    except Exception:  # noqa: BLE001 — fall back to a single (large) chunk
        return prefix, [inner], suffix
    chunks = [c for c in _chunk_container(wrapper, limit) if c.strip()]
    return prefix, (chunks or [inner]), suffix


async def _translate_chunks(api_key: str, chunks: list[str], language: str) -> tuple[str, str]:
    """Translate every chunk (concurrently, bounded) and join them back in order."""
    system = _translate_system_prompt(language)
    results: list[str] = [""] * len(chunks)
    model_used = DEFAULT_MODEL
    limiter = anyio.CapacityLimiter(TRANSLATE_MAX_CONCURRENCY)

    async def worker(idx: int, chunk: str) -> None:
        nonlocal model_used
        async with limiter:
            messages = [
                {"role": "system", "content": system},
                {"role": "user", "content": chunk},
            ]
            out, model = await anyio.to_thread.run_sync(lambda: call_openai(api_key, messages))
            results[idx] = _strip_code_fences(out)
            model_used = model

    async with anyio.create_task_group() as task_group:
        for idx, chunk in enumerate(chunks):
            task_group.start_soon(worker, idx, chunk)
    return "".join(results), model_used

_SUMMARY_STYLES = {
    "short": "3-5 sentences capturing only the essential points",
    "medium": "one or two well-structured paragraphs",
    "detailed": "a thorough summary with short sections and bullet points for key details",
}


async def _extract_upload(file: UploadFile) -> tuple[str, str, bool]:
    """Save + extract text; returns (text, original_name, truncated)."""
    workdir = new_workdir()
    try:
        src = await save_upload(file, workdir, SUPPORTED_EXTS)
        try:
            with anyio.fail_after(EXTRACTION_TIMEOUT_S):
                text = await anyio.to_thread.run_sync(extract_text, src, workdir, abandon_on_cancel=True)
        except TimeoutError:
            raise HTTPException(
                status_code=504,
                detail="This document took too long to process. Try a smaller or simpler file.",
            )
        return text, src.name, False
    finally:
        cleanup(workdir)


@router.post("/ai/summarize", dependencies=[require_permission("resources:write")])
async def summarize(
    file: UploadFile,
    length: str = Form(default="medium"),
    ctx: Any = _CTX_DEP,
):
    api_key = _require_openai_key(ctx)
    style = _SUMMARY_STYLES.get(length, _SUMMARY_STYLES["medium"])
    text, name, _ = await _extract_upload(file)
    truncated = len(text) > SUMMARIZE_MAX_CHARS
    if truncated:
        text = text[:SUMMARIZE_MAX_CHARS]

    messages = [
        {
            "role": "system",
            "content": (
                "You are an expert document analyst. Summarize the user's document as "
                f"{style}. Preserve factual accuracy, keep the original language of the "
                "document, and do not invent information. Output only the summary."
            ),
        },
        {"role": "user", "content": f"Document: {name}\n\n{text}"},
    ]
    summary, model = await anyio.to_thread.run_sync(
        lambda: call_openai(api_key, messages)
    )
    return {
        "summary": summary,
        "model": model,
        "filename": name,
        "characters_analyzed": len(text),
        "truncated": truncated,
    }


def _first_http_error(exc: BaseException) -> BaseException:
    """The first HTTPException inside a (possibly nested) exception group, or
    the exception itself when there is none."""
    if isinstance(exc, HTTPException):
        return exc
    for inner in getattr(exc, "exceptions", ()):
        found = _first_http_error(inner)
        if isinstance(found, HTTPException):
            return found
    return exc


@router.post("/ai/proofread", dependencies=[require_permission("resources:write")])
async def proofread(
    file: UploadFile,
    ctx: Any = _CTX_DEP,
):
    """Fix spelling/grammar INSIDE the document's own pages (see
    core/proofread.py), returning them in the editor's shape so the preview and
    the download keep the original layout and format."""
    from .edit import extract_pages

    api_key = _require_openai_key(ctx)
    workdir = new_workdir()
    try:
        src = await save_upload(file, workdir, EDIT_SUPPORTED_EXTS)
        with anyio.fail_after(EXTRACTION_TIMEOUT_S):
            pages, _ = await extract_pages(src, workdir)
        name = src.name
    except TimeoutError:
        raise HTTPException(
            status_code=504,
            detail="This document took too long to process. Try a smaller or simpler file.",
        )
    finally:
        cleanup(workdir)

    loaded, segments = proofreading.load_pages(pages)
    proofreading.check_size(segments)

    fixes: list[dict] = []
    model_used = DEFAULT_MODEL
    limiter = anyio.CapacityLimiter(proofreading.MAX_CONCURRENCY)

    async def worker(batch) -> None:
        nonlocal model_used
        async with limiter:
            messages = proofreading.messages_for(batch)
            content, model = await anyio.to_thread.run_sync(
                lambda: call_openai(api_key, messages, json_mode=True)
            )
            fixes.extend(proofreading.parse_fixes(content, {s.id for s in batch}))
            model_used = model

    try:
        async with anyio.create_task_group() as group:
            for batch in proofreading.batches(segments):
                group.start_soon(worker, batch)
    except Exception as exc:  # noqa: BLE001
        # A task group wraps a worker's HTTPException (bad key, rate limit,
        # unreadable reply) in an exception group; surface the real one so the
        # user sees its message, not a bare 500.
        raise _first_http_error(exc) from None

    # Document order, so the list reads top to bottom however batches finished.
    order = {s.id: i for i, s in enumerate(segments)}
    fixes.sort(key=lambda f: order[f["id"]])
    applied = proofreading.apply_fixes(loaded, segments, fixes)
    return {
        "pages": proofreading.dump_pages(loaded),
        "corrected_text": proofreading.corrected_text(segments),
        "fixes": applied,
        "fix_count": sum(1 for f in applied if f["applied"]),
        "model": model_used,
        "filename": name,
    }


@router.post("/ai/translate", dependencies=[require_permission("resources:write")])
async def translate(
    file: UploadFile,
    target_language: str = Form(...),
    ctx: Any = _CTX_DEP,
):
    api_key = _require_openai_key(ctx)
    workdir = new_workdir()
    try:
        src = await save_upload(file, workdir, EDIT_SUPPORTED_EXTS)
        html = await extract_html_bounded(src, workdir)
        name = src.name
    finally:
        cleanup(workdir)

    if not html.strip():
        raise HTTPException(status_code=422, detail="No readable content found in this document.")

    # Strip embedded images before the round-trip to the LLM: they can be
    # megabytes of token-expensive base64 the model doesn't need to see (and
    # could corrupt); placeholders are swapped back in afterwards.
    stripped, image_map = extract_data_images(html)

    language = (target_language or "").strip() or "English"
    # Always translate the WHOLE document: split it into element-boundary chunks
    # (so tables/rows/cells are never cut) and translate every chunk.
    prefix, chunks, suffix = _split_html_chunks(stripped, TRANSLATE_CHUNK_CHARS)
    if chunks:
        translated_inner, model = await _translate_chunks(api_key, chunks, language)
    else:
        translated_inner, model = "", DEFAULT_MODEL

    translated = restore_data_images(prefix + translated_inner + suffix, image_map)

    return {
        "html": translated,
        "model": model,
        "filename": name,
        "target_language": language,
        "truncated": False,
    }


@router.post("/ai/export", dependencies=[require_permission("resources:write")])
async def export_text(
    text: str = Form(...),
    format: str = Form(default="txt"),
    basename: str = Form(default="result"),
):
    if format not in {"txt", "pdf", "docx"}:
        raise HTTPException(status_code=422, detail="Export format must be 'txt', 'pdf' or 'docx'.")
    if not text.strip():
        raise HTTPException(status_code=422, detail="Nothing to export.")

    stem = safe_name(basename, "result").rsplit(".", 1)[0] or "result"
    workdir = new_workdir()
    try:
        if format == "txt":
            out = workdir / f"{stem}.txt"
            out.write_text(text, encoding="utf-8")
            return respond_with(workdir, [out])
        if format == "docx":
            from docx import Document

            document = Document()
            for paragraph in text.split("\n"):
                document.add_paragraph(paragraph)
            out = workdir / f"{stem}.docx"
            document.save(str(out))
            return respond_with(workdir, [out])
        # pdf — reuse the Text→PDF converter on a temp .txt source
        from ..core.convert_ops import text_to_pdf

        src = workdir / f"{stem}.txt"
        src.write_text(text, encoding="utf-8")
        outputs = await anyio.to_thread.run_sync(text_to_pdf, src, workdir)
        return respond_with(workdir, outputs)
    except HTTPException:
        cleanup(workdir)
        raise
    except Exception as exc:  # noqa: BLE001
        cleanup(workdir)
        raise HTTPException(status_code=500, detail=f"Export failed: {exc}")


# ------------------------------------------------------------ Edit assistant

class AssistantRequest(BaseModel):
    """What the editor's AI panel sends: the open document as an outline (see
    frontend/lib/doc-assistant.ts), the conversation so far, and the prompt."""

    kind: str = Field(default="document", pattern="^(document|sheet)$")
    layout: str = "flowing"
    blocks: list[dict] = Field(default_factory=list)
    cells: list[dict] = Field(default_factory=list)
    objects: list[dict] = Field(default_factory=list)
    selection: str = Field(default="", max_length=4000)
    history: list[dict] = Field(default_factory=list)
    prompt: str = Field(min_length=1, max_length=4000)


@router.post("/ai/assistant", dependencies=[require_permission("resources:write")])
async def edit_assistant(body: AssistantRequest, ctx: Any = _CTX_DEP):
    api_key = _require_openai_key(ctx)
    messages = assistant.build_messages(
        kind=body.kind,
        blocks=body.blocks,
        cells=body.cells,
        layout=body.layout,
        history=body.history,
        prompt=body.prompt,
        objects=body.objects,
        selection=body.selection,
    )
    content, model = await anyio.to_thread.run_sync(
        lambda: call_openai(api_key, messages, json_mode=True)
    )
    reply = assistant.parse_reply(
        content,
        kind=body.kind,
        blocks=body.blocks,
        cells=body.cells,
        objects=body.objects,
        selection=body.selection,
    )
    return {**reply, "model": model}


# ------------------------------------------------------------ Document chat
#
# The AI Summarizer is a conversation about one document: summaries of any
# length are just suggested prompts, and any other question works the same
# way. The document is read ONCE (/ai/chat/open) and its text travels with each
# message, so the server keeps no session and nothing is re-extracted per turn.

CHAT_HISTORY_TURNS = 10


@router.post("/ai/chat/open", dependencies=[require_permission("resources:write")])
async def open_document_chat(file: UploadFile, ctx: Any = _CTX_DEP):
    _require_openai_key(ctx)
    text, name, _ = await _extract_upload(file)
    if not text.strip():
        raise HTTPException(status_code=422, detail="No readable text was found in this document.")
    truncated = len(text) > SUMMARIZE_MAX_CHARS
    return {
        "text": text[:SUMMARIZE_MAX_CHARS],
        "filename": name,
        "characters": len(text),
        "truncated": truncated,
    }


class DocumentChatRequest(BaseModel):
    filename: str = "document"
    text: str = Field(min_length=1, max_length=SUMMARIZE_MAX_CHARS + 1000)
    history: list[dict] = Field(default_factory=list)
    prompt: str = Field(min_length=1, max_length=4000)
    #: A summary chip's length ("short" | "medium" | "detailed"), so the
    #: suggested summaries keep the exact shapes the old length option had.
    summary: str | None = None


@router.post("/ai/chat", dependencies=[require_permission("resources:write")])
async def document_chat(body: DocumentChatRequest, ctx: Any = _CTX_DEP):
    api_key = _require_openai_key(ctx)
    system = (
        "You are a helpful assistant answering questions about ONE document, "
        "given below. Base every answer on the document; if it does not contain "
        "the answer, say so plainly instead of guessing. Quote or point to the "
        "relevant part when that helps. Reply in the language the user writes in. "
        "Use plain text: short paragraphs, and '- ' bullets for lists (no Markdown "
        "headings or bold). Do exact arithmetic carefully and show the figures used.\n\n"
        f"Document: {body.filename}\n<<<\n{body.text}\n>>>"
    )
    messages: list[dict] = [{"role": "system", "content": system}]
    for turn in body.history[-CHAT_HISTORY_TURNS:]:
        role, content = turn.get("role"), str(turn.get("content", ""))[:8000]
        if role in ("user", "assistant") and content:
            messages.append({"role": role, "content": content})
    prompt = body.prompt
    style = _SUMMARY_STYLES.get(body.summary or "")
    if style:
        prompt += f"\n\n(Write the summary as {style}. Preserve factual accuracy and do not invent anything.)"
    messages.append({"role": "user", "content": prompt})
    answer, model = await anyio.to_thread.run_sync(lambda: call_openai(api_key, messages))
    return {"answer": answer, "model": model}


# -------------------------------------------------------------- Global chat

class AgentRequest(BaseModel):
    """The global chat's turn: the tool catalog (from the frontend registry),
    the session's files by id, the conversation, and the new message."""

    catalog: list[dict] = Field(default_factory=list, max_length=200)
    files: list[dict] = Field(default_factory=list, max_length=200)
    history: list[dict] = Field(default_factory=list)
    prompt: str = Field(min_length=1, max_length=4000)


@router.post("/ai/agent", dependencies=[require_permission("resources:write")])
async def chat_agent_turn(body: AgentRequest, ctx: Any = _CTX_DEP):
    api_key = _require_openai_key(ctx)
    messages = chat_agent.build_messages(
        catalog=body.catalog, files=body.files, history=body.history, prompt=body.prompt
    )
    content, model = await anyio.to_thread.run_sync(
        lambda: call_openai(api_key, messages, json_mode=True)
    )
    plan = chat_agent.parse_plan(content, catalog=body.catalog, files=body.files)
    return {**plan, "model": model}
