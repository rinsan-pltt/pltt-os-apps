"""CSV / XLSX / PDF builders for reimbursement reports.

Kept separate from `routes/reports.py` (which only handles the HTTP layer
and DB query), same core-vs-route split as the reference file-convertor
plugin's `core/*_ops.py` modules.
"""

from __future__ import annotations

import csv
import io

from .models import Expense


# Shared by the CSV and XLSX writers so the two exports can never drift apart.
_SHEET_COLUMNS = ["Date", "Vendor", "Category", "Amount", "Currency", "Status", "Notes"]


def build_csv(expenses: list[Expense]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(_SHEET_COLUMNS)
    for e in expenses:
        writer.writerow(
            [
                e.expense_date.isoformat(),
                e.vendor,
                e.category,
                f"{e.amount:.2f}",
                e.currency,
                e.status,
                e.notes or "",
            ]
        )
    total = sum(e.amount for e in expenses)
    writer.writerow([])
    writer.writerow(["", "", "", f"{total:.2f}", "", "TOTAL", ""])
    return buffer.getvalue().encode("utf-8")



def build_xlsx(expenses: list[Expense], title: str) -> bytes:
    """A real Office spreadsheet, not a CSV with an .xlsx name.

    The point of offering this next to the CSV is the things a CSV cannot carry:
    Amount lands as a NUMBER with a currency-style format rather than a string,
    the date is a real date cell, and the header row is frozen so a long report
    stays readable while scrolling. A renamed CSV would open with a corrupt-file
    warning in Excel, so this uses openpyxl — already a declared dependency,
    used on the import side.
    """
    # Imported here rather than at module scope: the CSV and PDF paths are the
    # common ones and shouldn't pay for loading openpyxl.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    # Excel rejects sheet names over 31 chars and the characters below, and the
    # title is caller-supplied.
    safe = "".join(c for c in (title or "Report") if c not in "[]:*?/\\")[:31] or "Report"
    sheet.title = safe

    sheet.append(_SHEET_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"

    for e in expenses:
        sheet.append(
            [
                e.expense_date,
                e.vendor,
                e.category,
                float(e.amount),
                e.currency,
                e.status,
                e.notes or "",
            ]
        )

    total_row = sheet.max_row + 2
    sheet.cell(row=total_row, column=3, value="TOTAL").font = Font(bold=True)
    total_cell = sheet.cell(row=total_row, column=4, value=float(sum(e.amount for e in expenses)))
    total_cell.font = Font(bold=True)

    # Numbers as numbers, dates as dates — the reason to prefer this over CSV.
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "#,##0.00"

    # Width from the widest cell in each column, clamped: a stray long note
    # should not push the Notes column off the screen.
    for index, header in enumerate(_SHEET_COLUMNS, start=1):
        longest = max(
            [len(header)]
            + [len(str(sheet.cell(row=r, column=index).value or "")) for r in range(2, sheet.max_row + 1)]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 44)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# (header label, column width as a fraction of the usable page width, alignment).
# Fractions sum to 1.0. Amount is right-aligned so the decimals line up.
_COLUMNS = [
    ("Date", 0.12, "left"),
    ("Vendor", 0.26, "left"),
    ("Category", 0.17, "left"),
    ("Amount", 0.17, "right"),
    ("Status", 0.13, "left"),
    ("Notes", 0.15, "left"),
]

_CELL_PAD = 4.0  # horizontal padding inside each cell, in points
_HEADER_FONT = "hebo"  # Helvetica-Bold
_BODY_FONT = "helv"  # Helvetica
_HEADER_SIZE = 9.0
_BODY_SIZE = 8.5
_GRID_COLOR = (0.80, 0.82, 0.82)
_HEADER_FILL = (0.93, 0.95, 0.94)


def build_pdf(expenses: list[Expense], title: str) -> bytes:
    import fitz

    doc = fitz.open()
    rect = fitz.paper_rect("a4")
    margin = 40.0
    usable_width = rect.width - 2 * margin
    row_height = 18.0
    header_height = 20.0
    bottom_limit = rect.height - margin
    total = sum(e.amount for e in expenses)

    # Precompute each column's left edge (x0) and width in absolute points.
    cols: list[tuple[str, float, float, str]] = []
    x = margin
    for header, frac, align in _COLUMNS:
        width = usable_width * frac
        cols.append((header, x, width, align))
        x += width
    # Vertical grid-line positions: left margin + every column's right edge.
    boundaries = [margin] + [x0 + w for _, x0, w, _ in cols]

    def fit_text(text: str, max_width: float, fontname: str, fontsize: float) -> str:
        """Clip `text` to `max_width` points, appending an ellipsis if trimmed,
        so a value never spills into the neighbouring column."""
        text = " ".join(str(text).split())  # collapse newlines/runs of spaces
        if max_width <= 0 or not text:
            return ""
        if fitz.get_text_length(text, fontname=fontname, fontsize=fontsize) <= max_width:
            return text
        # ASCII ellipsis — PyMuPDF's built-in Helvetica renders the Unicode "…"
        # as a bare dot, so "..." reads unambiguously as truncation.
        ellipsis = "..."
        while text and fitz.get_text_length(text + ellipsis, fontname=fontname, fontsize=fontsize) > max_width:
            text = text[:-1]
        return (text + ellipsis) if text else ellipsis

    def draw_cell(page, text, x0, width, baseline, fontname, fontsize, align) -> None:
        fitted = fit_text(text, width - 2 * _CELL_PAD, fontname, fontsize)
        if align == "right":
            tw = fitz.get_text_length(fitted, fontname=fontname, fontsize=fontsize)
            tx = x0 + width - _CELL_PAD - tw
        else:
            tx = x0 + _CELL_PAD
        page.insert_text((tx, baseline), fitted, fontsize=fontsize, fontname=fontname)

    def draw_grid_row(page, y_top, height) -> None:
        """Vertical column separators for this row + its bottom border."""
        for xb in boundaries:
            page.draw_line((xb, y_top), (xb, y_top + height), color=_GRID_COLOR, width=0.5)
        page.draw_line((margin, y_top + height), (margin + usable_width, y_top + height), color=_GRID_COLOR, width=0.5)

    def start_page() -> tuple["fitz.Page", float]:
        page = doc.new_page(width=rect.width, height=rect.height)
        y = margin
        page.insert_text((margin, y + 14), title, fontsize=16, fontname=_HEADER_FONT)
        y += 26
        page.insert_text(
            (margin, y + 6),
            f"{len(expenses)} expense(s)  ·  Total {total:.2f}",
            fontsize=9,
            fontname=_BODY_FONT,
            color=(0.4, 0.4, 0.4),
        )
        y += 18
        # Header row: shaded band, top border, labels, then the grid.
        page.draw_rect(
            fitz.Rect(margin, y, margin + usable_width, y + header_height),
            color=None,
            fill=_HEADER_FILL,
        )
        page.draw_line((margin, y), (margin + usable_width, y), color=_GRID_COLOR, width=0.5)
        baseline = y + header_height - 6
        for header, x0, width, align in cols:
            draw_cell(page, header, x0, width, baseline, _HEADER_FONT, _HEADER_SIZE, align)
        draw_grid_row(page, y, header_height)
        return page, y + header_height

    page, y = start_page()
    for e in expenses:
        if y + row_height > bottom_limit:
            page, y = start_page()
        baseline = y + row_height - 6
        values = [
            e.expense_date.isoformat(),
            e.vendor,
            e.category,
            f"{e.currency} {e.amount:.2f}",
            e.status,
            e.notes or "",
        ]
        for value, (_, x0, width, align) in zip(values, cols):
            draw_cell(page, value, x0, width, baseline, _BODY_FONT, _BODY_SIZE, align)
        draw_grid_row(page, y, row_height)
        y += row_height

    # Total row: heavier rule above, "TOTAL" under Category, sum under Amount.
    if y + row_height > bottom_limit:
        page, y = start_page()
    page.draw_line((margin, y), (margin + usable_width, y), color=(0.55, 0.57, 0.57), width=1.0)
    baseline = y + row_height - 6
    _, cat_x0, cat_w, _ = cols[2]
    _, amt_x0, amt_w, _ = cols[3]
    draw_cell(page, "TOTAL", cat_x0, cat_w, baseline, _HEADER_FONT, _HEADER_SIZE, "left")
    draw_cell(page, f"{total:.2f}", amt_x0, amt_w, baseline, _HEADER_FONT, _HEADER_SIZE, "right")
    draw_grid_row(page, y, row_height)

    out = doc.tobytes()
    doc.close()
    return out
